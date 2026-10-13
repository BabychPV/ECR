"""WW-02: cell-формули (рівня комірки) з аркушів 7. Water Report, 7a, 7b, 8. Waste Report -> секція `cellFormulas` моделі.

Методика та сама, що для 7.0 (gen_ecr_model._cell_formulas_7_0, LP-31):
  * Excel рахує порожню клітинку як 0, а `a+b` в ECR при відсутньому операнді дає null (ТЗ 05-engines «Семантика null»),
    тому кожна лінійна формула пишеться в null-безпечній формі: `SUM(плюси)`, `SUM(плюси)-SUM(мінуси)`, з масштабом
    `/k` або `*k` над усім виразом. SUM поглинає null, порожня множина = 0 (02b §6.1).
  * Лише доведені класи: сума/різниця комірок і діапазонів (по одному стовпцю чи рядку), дужки, `SUM(...)`, унарний
    мінус, масштаб числом над ВСІМ виразом. Усе інше (IF, VLOOKUP, текст, посилання поза моделлю, нечислові колонки) -
    НЕ вгадується: потрапляє в `cellFormulasUnresolved` з причиною або в `cellFormulasSkipped` (довідникові підписи).
  * Рядок із комірочними формулами позначається `readOnly` (Build -> RowDef.IsReadOnly): колонка лишається ручного вводу
    для решти рядків, а рядок підсумку/балансу людина не редагує (ECR-TMPL-4227 для Cell дозволений лише на
    обчислюваній колонці АБО в рядку лише для читання).
"""
import re
import random
from openpyxl.utils import column_index_from_string, get_column_letter

MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November',
          'December']
SHEET_PREFIXES = ('7. ', '7a', '7b', '8. ')
LABEL_FN = re.compile(r'VLOOKUP|IFERROR|TEXTBEFORE|TEXTAFTER|INDEX\(|MATCH\(', re.I)
TOKEN = re.compile(
    r"\s*(?:(?P<num>\d+(?:\.\d+)?)"
    r"|(?P<fn>[A-Za-z_][A-Za-z0-9_.]*)\("
    r"|(?P<sheet>'[^']+'!)?(?P<ref>\$?[A-Z]{1,3}\$?\d+)(?::(?P<ref2>\$?[A-Z]{1,3}\$?\d+))?"
    r"|(?P<op>[-+*/(),]))")


class Unresolved(Exception):
    pass


def _tokenize(f):
    pos, out = 0, []
    while pos < len(f):
        if f[pos].isspace():
            pos += 1
            continue
        m = TOKEN.match(f, pos)
        if not m or m.end() == pos:
            raise Unresolved('unsupported token near "%s"' % f[pos:pos + 12])
        pos = m.end()
        if m.group('num') is not None:
            out.append(('num', float(m.group('num'))))
        elif m.group('fn') is not None:
            out.append(('fn', m.group('fn').upper()))
        elif m.group('ref') is not None:
            sh = m.group('sheet')
            out.append(('ref', sh[1:-2] if sh else None, m.group('ref').replace('$', ''),
                        m.group('ref2').replace('$', '') if m.group('ref2') else None))
        else:
            out.append(('op', m.group('op')))
    return out


class _Parser:
    def __init__(self, toks):
        self.t, self.i = toks, 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self):
        x = self.t[self.i]
        self.i += 1
        return x

    def expr(self):
        items = [(1, self.term())]
        while (p := self.peek()) and p[0] == 'op' and p[1] in '+-':
            self.take()
            items.append((1 if p[1] == '+' else -1, self.term()))
        return ('add', items) if len(items) > 1 else items[0][1]

    def term(self):
        first = self.factor()
        scales = []
        while (p := self.peek()) and p[0] == 'op' and p[1] in '*/':
            self.take()
            scales.append((p[1], self.factor()))
        return ('mul', first, scales) if scales else first

    def factor(self):
        p = self.peek()
        if p is None:
            raise Unresolved('unexpected end of formula')
        if p == ('op', '-'):
            self.take()
            return ('add', [(-1, self.factor())])
        if p == ('op', '+'):
            self.take()
            return self.factor()
        if p == ('op', '('):
            self.take()
            e = self.expr()
            if self.peek() != ('op', ')'):
                raise Unresolved('unbalanced parentheses')
            self.take()
            return e
        if p[0] == 'num':
            self.take()
            return ('num', p[1])
        if p[0] == 'ref':
            self.take()
            return ('range', p[1], p[2], p[3]) if p[3] else ('cell', p[1], p[2])
        if p[0] == 'fn':
            self.take()
            if p[1] != 'SUM':
                raise Unresolved('function %s is not translated' % p[1])
            args = []
            if self.peek() == ('op', ')'):
                self.take()
                return ('sum', args)
            while True:
                args.append(self.expr())
                q = self.take()
                if q == ('op', ')'):
                    break
                if q != ('op', ','):
                    raise Unresolved('bad SUM argument list')
            return ('sum', args)
        raise Unresolved('unexpected token %s' % (p,))


def _split(a1):
    m = re.match(r'^([A-Z]+)(\d+)$', a1)
    return m.group(1), int(m.group(2))


def _flatten(node, sign, out):
    """Лінійна форма: список (знак, (аркуш|None, колонка, рядок))."""
    k = node[0]
    if k == 'add':
        for s, ch in node[1]:
            _flatten(ch, sign * s, out)
    elif k == 'cell':
        c, r = _split(node[2])
        out.append((sign, (node[1], c, r)))
    elif k == 'range':
        c1, r1 = _split(node[2])
        c2, r2 = _split(node[3])
        if c1 != c2 and r1 != r2:
            raise Unresolved('two-dimensional range %s:%s' % (node[2], node[3]))
        if c1 == c2:
            for r in range(r1, r2 + 1):
                out.append((sign, (node[1], c1, r)))
        else:
            for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                out.append((sign, (node[1], get_column_letter(ci), r1)))
    elif k == 'sum':
        for a in node[1]:
            _flatten(a, sign, out)
    elif k == 'num':
        raise Unresolved('numeric constant inside a sum')
    else:
        raise Unresolved('nested scaling/product is not a linear sum')


def linearize(formula):
    """-> (terms[(sign, (sheet, col, row))], scale[(op, num)] ). Unresolved при не доведеному класі."""
    toks = _tokenize(formula)
    p = _Parser(toks)
    tree = p.expr()
    if p.peek() is not None:
        raise Unresolved('trailing tokens')
    scale = []
    if tree[0] == 'mul':
        _, base, ss = tree
        for op, f in ss:
            if f[0] != 'num':
                raise Unresolved('product/quotient of cells')
            scale.append((op, f[1]))
        tree = base
    terms = []
    _flatten(tree, 1, terms)
    if not terms:
        raise Unresolved('no cell operands')
    return terms, scale


def _num(v):
    return ('%g' % v)


def build_expression(refs, signs, scale):
    plus = [r for r, s in zip(refs, signs) if s > 0]
    minus = [r for r, s in zip(refs, signs) if s < 0]
    if plus and minus:
        body = 'SUM(%s)-SUM(%s)' % (','.join(plus), ','.join(minus))
    elif plus:
        body = 'SUM(%s)' % ','.join(plus)
    else:
        body = '-SUM(%s)' % ','.join(minus)
    for op, k in scale:
        body = '%s%s%s' % (('(%s)' % body) if (plus and minus) or (not plus) else body, op, _num(k))
    return body


# ---- незалежна перевірка семантики Excel: порожнє = 0 (власна інтерпретація формули, не через linearize) --------
def _excel_eval(formula, values):
    def flat(*a):
        t = 0.0
        for x in a:
            t += sum(x) if isinstance(x, list) else x
        return t

    def sheetref(m):
        c, r = _split(m.group('ref').replace('$', ''))
        return 'V(%r,%r,%d)' % (m.group('sheet')[1:-2] if m.group('sheet') else None, c, r)

    py = re.sub(
        r"(?P<sheet>'[^']+'!)?(?P<ref>\$?[A-Z]{1,3}\$?\d+):(?P<ref2>\$?[A-Z]{1,3}\$?\d+)",
        lambda m: 'RNG(%r,%r,%r)' % (m.group('sheet')[1:-2] if m.group('sheet') else None,
                                      m.group('ref').replace('$', ''), m.group('ref2').replace('$', '')), formula)
    py = re.sub(r"(?<![A-Za-z'])(?P<sheet>'[^']+'!)?(?P<ref>\$?[A-Z]{1,3}\$?\d+)(?![\d(])", sheetref, py)
    py = py.replace('SUM(', 'S(')

    def V(sh, c, r):
        return values.get((sh, c, r), 0.0)

    def RNG(sh, a, b):
        c1, r1 = _split(a)
        c2, r2 = _split(b)
        return [values.get((sh, get_column_letter(ci), r), 0.0)
                for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1)
                for r in range(r1, r2 + 1)]

    return eval(py, {'__builtins__': {}}, {'V': V, 'RNG': RNG, 'S': flat})


def apply(out, PS):
    stat = {'sheets': {}, 'total': 0, 'unresolved': 0, 'skipped': 0}
    rnd = random.Random(20261013)
    for sh in out['sheets']:
        name = sh['excelName']
        if not name.startswith(SHEET_PREFIXES):
            continue
        cells = PS[name]['cells']
        # excelRow -> (table, row); дублікати Group (T8_00/T8_01) віддають перевагу першій таблиці.
        row_index = {}
        for t in sh['tables']:
            for rw in t['rows']:
                row_index.setdefault(rw['excelRow'], (t, rw))
        col_letter = {t['code']: {c['excel']: c for c in t['columns'] if c.get('excel')} for t in sh['tables']}
        results = {t['code']: [] for t in sh['tables']}
        unres = {t['code']: [] for t in sh['tables']}
        skipped = {t['code']: {} for t in sh['tables']}
        readonly_rows = {}
        for addr, x in cells.items():
            f = x.get('f')
            if not f:
                continue
            L, r = _split(addr)
            tgt = row_index.get(r)
            t0 = tgt[0] if tgt else sh['tables'][0]

            def bad(why, t=t0, addr=addr, f=f):
                unres[t['code']].append({'cell': addr, 'formula': f, 'why': why})

            if LABEL_FN.search(f):
                cls = 'label/lookup formula (Configuration/header), not a numeric calculation'
                s = skipped[t0['code']].setdefault(cls, {'class': cls, 'count': 0, 'sample': '%s: %s' % (addr, f[:90])})
                s['count'] += 1
                continue
            if tgt is None:
                bad('row %d is not a row of any table of the model' % r)
                continue
            t, rw = tgt
            col = col_letter[t['code']].get(L)
            if col is None:
                bad('column %s is not in the model table %s' % (L, t['code']))
                continue
            if col['dataType'] not in ('Decimal', 'Int', 'Formula'):
                bad('target column %s (%s) is not numeric: %s' % (L, col['code'], col['dataType']))
                continue
            try:
                terms, scale = linearize(f)
                refs, signs, problems = [], [], []
                tcode = col['code']
                for s, (rsheet, rc, rr) in terms:
                    src_sheet = rsheet or name
                    src = next((q for q in out['sheets'] if q['excelName'] == src_sheet), None)
                    if src is None:
                        raise Unresolved('sheet %s is not in the model' % src_sheet)
                    stbl_rows = None
                    for st in src['tables']:
                        hit = next((w for w in st['rows'] if w['excelRow'] == rr), None)
                        if hit and (stbl_rows is None or hit['kind'] != 'Group'):
                            stbl_rows = (st, hit)
                            break
                    if stbl_rows is None:
                        raise Unresolved('row %d of %s is not a row of the model' % (rr, src_sheet))
                    st, hit = stbl_rows
                    c2 = next((c for c in st['columns'] if c.get('excel') == rc), None)
                    if c2 is None:
                        raise Unresolved('column %s of %s is not in the model' % (rc, src_sheet))
                    if c2['dataType'] not in ('Decimal', 'Int', 'Formula'):
                        raise Unresolved('operand %s%d is not numeric (%s)' % (rc, rr, c2['dataType']))
                    # Колонка джерела: та сама, що цільова; лише в тому ж аркуші допустимий «попередній місяць» /
                    # OpeningBalance (баланс Water Report).
                    ok = c2['code'] == tcode
                    if not ok and src_sheet == name and tcode in MONTHS:
                        mi = MONTHS.index(tcode)
                        ok = (mi > 0 and c2['code'] == MONTHS[mi - 1]) or (mi == 0 and c2['code'] == 'OpeningBalance')
                    if not ok:
                        raise Unresolved('operand %s%d is in column %s, not in %s' % (rc, rr, c2['code'], tcode))
                    refs.append('[%s].[%s].[%s].[%s]' % (src['code'], st['code'], hit['key'], c2['code']))
                    signs.append(s)
                expr = build_expression(refs, signs, scale)
            except Unresolved as e:
                bad(str(e))
                continue
            # Незалежна перевірка: Excel-інтерпретація формули (порожнє = 0) == знакова сума, на випадкових входах.
            keys = {(rs, c, rr) for _, (rs, c, rr) in terms}
            for _ in range(3):
                vals = {k: float(rnd.randint(1, 9000)) for k in keys}
                want = _excel_eval(f, vals)
                got = sum(s * vals[k] for s, k in terms)
                for op, kf in scale:
                    got = got * kf if op == '*' else got / kf
                assert abs(want - got) < 1e-6 * max(1.0, abs(want)), 'self-check %s: %s vs %s' % (addr, want, got)
            results[t['code']].append({
                'row': rw['key'], 'column': tcode, 'expression': expr, 'excelCell': addr, 'excelFormula': f,
                'terms': ['%s%s' % ('+' if s > 0 else '-', r_) for s, r_ in zip(signs, refs)],
                'scale': ['%s%s' % (op, _num(k)) for op, k in scale]})
            readonly_rows.setdefault(t['code'], set()).add(rw['key'])
        for t in sh['tables']:
            c = t['code']
            res = sorted(results[c], key=lambda q: (q['row'], MONTHS.index(q['column']) if q['column'] in MONTHS else 99,
                                                    q['column']))
            t['cellFormulas'] = res
            t['cellFormulasUnresolved'] = sorted(unres[c], key=lambda q: (_split(q['cell'])[1], q['cell']))
            t['cellFormulasSkipped'] = list(skipped[c].values())
            ro = readonly_rows.get(c, set())
            for rw in t['rows']:
                if rw['key'] in ro:
                    rw['readOnly'] = True
            stat['total'] += len(res)
            stat['unresolved'] += len(t['cellFormulasUnresolved'])
            stat['skipped'] += sum(s['count'] for s in t['cellFormulasSkipped'])
            print('WW-02 cellFormulas %s/%s: %d, unresolved %d, skipped %d, readOnly rows %d' % (
                sh['code'], c, len(res), len(t['cellFormulasUnresolved']),
                sum(s['count'] for s in t['cellFormulasSkipped']), len(ro)))
    return stat
