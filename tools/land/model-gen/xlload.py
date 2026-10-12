import sys, re, json, pickle, zipfile, colorsys
import xml.etree.ElementTree as ET
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'x14': 'http://schemas.microsoft.com/office/spreadsheetml/2009/9/main',
      'xm': 'http://schemas.microsoft.com/office/excel/2006/main'}
M = '{%s}' % NS['m']
R = '{%s}' % NS['r']
X14 = '{%s}' % NS['x14']
XM = '{%s}' % NS['xm']

cell_re = re.compile(r'^([A-Z]+)(\d+)$')


def split_ref(ref):
    m = cell_re.match(ref)
    return m.group(1), int(m.group(2))


def load_shared(z):
    out = []
    root = ET.fromstring(z.read('xl/sharedStrings.xml'))
    for si in root.findall(M + 'si'):
        out.append(''.join(t.text or '' for t in si.iter(M + 't')))
    return out


def load_workbook_meta(z):
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    relmap = {r.get('Id'): r.get('Target') for r in rels}
    sheets = []
    for s in wb.find(M + 'sheets'):
        rid = s.get(R + 'id')
        t = relmap[rid]
        t = t.lstrip('/')
        if not t.startswith('xl/'):
            t = 'xl/' + t
        sheets.append({'name': s.get('name'), 'sheetId': s.get('sheetId'), 'state': s.get('state', 'visible'), 'path': t})
    names = []
    dn = wb.find(M + 'definedNames')
    if dn is not None:
        for d in dn:
            names.append({'name': d.get('name'), 'localSheetId': d.get('localSheetId'), 'hidden': d.get('hidden'), 'text': d.text})
    return sheets, names

# Стандартна індексована палітра Excel (у книзі немає custom indexedColors — перевірено: <colors> містить лише mruColors).
INDEXED = ['000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
           '000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
           '800000', '008000', '000080', '808000', '800080', '008080', 'C0C0C0', '808080',
           '9999FF', '993366', 'FFFFCC', 'CCFFFF', '660066', 'FF8080', '0066CC', 'CCCCFF',
           '000080', 'FF00FF', 'FFFF00', '00FFFF', '800080', '800000', '008080', '0000FF',
           '00CCFF', 'CCFFFF', 'CCFFCC', 'FFFF99', '99CCFF', 'FF99CC', 'CC99FF', 'FFCC99',
           '3366FF', '33CCCC', '99CC00', 'FFCC00', 'FF9900', 'FF6600', '666699', '969696',
           '003366', '339966', '003300', '333300', '993300', '993366', '333399', '333333',
           '000000', 'FFFFFF']
# індекс теми Excel -> позиція в clrScheme (dk1, lt1, dk2, lt2, accent1..6, hlink, folHlink)
THEME_POS = [1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11]
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
BUILTIN_NUMFMT = {1: '0', 2: '0.00', 3: '#,##0', 4: '#,##0.00', 9: '0%', 10: '0.00%', 11: '0.00E+00', 12: '# ?/?', 13: '# ??/??',
                  37: '#,##0 ;(#,##0)', 38: '#,##0 ;[Red](#,##0)', 39: '#,##0.00;(#,##0.00)', 40: '#,##0.00;[Red](#,##0.00)',
                  45: 'mm:ss', 46: '[h]:mm:ss', 47: 'mmss.0', 48: '##0.0E+0', 49: '@'}


def load_theme(z):
    try:
        root = ET.fromstring(z.read('xl/theme/theme1.xml'))
    except KeyError:
        return None
    cs = root.find('.//' + A + 'clrScheme')
    out = []
    for ch in cs:
        c = ch[0]
        out.append((c.get('val') if c.tag == A + 'srgbClr' else c.get('lastClr')) or '000000')
    return out


def tint_rgb(hexrgb, tint):
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def color_hex(el, theme):
    """<color/fgColor/bgColor> -> 'RRGGBB' або None (auto/відсутній)."""
    if el is None:
        return None
    if el.get('rgb'):
        return el.get('rgb')[-6:].upper()
    if el.get('indexed') is not None:
        i = int(el.get('indexed'))
        return INDEXED[i] if i < len(INDEXED) else None
    if el.get('theme') is not None and theme:
        i = int(el.get('theme'))
        base = theme[THEME_POS[i]] if i < len(THEME_POS) else None
        if base is None:
            return None
        t = float(el.get('tint') or 0)
        return tint_rgb(base.upper(), t) if t else base.upper()
    return None


def parse_styles(z):
    """xl/styles.xml -> {'xfs': [...], 'dxfs': [...]}: розкладені xf (заливка, шрифт, формат числа, захист) і dxf умовних форматів."""
    theme = load_theme(z)
    st = ET.fromstring(z.read('xl/styles.xml'))
    numfmts = {}
    nf = st.find(M + 'numFmts')
    if nf is not None:
        for n in nf:
            numfmts[int(n.get('numFmtId'))] = n.get('formatCode')
    fills = []
    for f in st.find(M + 'fills'):
        pf = f.find(M + 'patternFill')
        if pf is None:
            fills.append({'pattern': None, 'fg': None, 'bg': None})
            continue
        fills.append({'pattern': pf.get('patternType'), 'fg': color_hex(pf.find(M + 'fgColor'), theme), 'bg': color_hex(pf.find(M + 'bgColor'), theme)})
    fonts = []
    for f in st.find(M + 'fonts'):
        b = f.find(M + 'b')
        i = f.find(M + 'i')
        sz = f.find(M + 'sz')
        nm = f.find(M + 'name')
        fonts.append({'bold': b is not None and b.get('val') not in ('0', 'false'), 'italic': i is not None and i.get('val') not in ('0', 'false'),
                      'color': color_hex(f.find(M + 'color'), theme), 'name': nm.get('val') if nm is not None else None,
                      'size': float(sz.get('val')) if sz is not None else None})
    borders = []
    for b in st.find(M + 'borders'):
        borders.append({s: (b.find(M + s).get('style') if b.find(M + s) is not None else None) for s in ('left', 'right', 'top', 'bottom')})
    xfs = []
    for xf in st.find(M + 'cellXfs'):
        fid = int(xf.get('numFmtId') or 0)
        al = xf.find(M + 'alignment')
        pr = xf.find(M + 'protection')
        fill = fills[int(xf.get('fillId') or 0)]
        font = fonts[int(xf.get('fontId') or 0)]
        xfs.append({'fill': fill['fg'] if fill['pattern'] == 'solid' else None, 'pattern': fill['pattern'] if fill['pattern'] not in (None, 'none', 'solid') else None,
                    'fontColor': font['color'], 'bold': font['bold'], 'italic': font['italic'], 'fontName': font['name'], 'fontSize': font['size'],
                    'numFmt': None if fid == 0 else (numfmts.get(fid) or BUILTIN_NUMFMT.get(fid) or 'builtin:%d' % fid),
                    'wrap': al is not None and al.get('wrapText') in ('1', 'true'), 'halign': al.get('horizontal') if al is not None else None,
                    'valign': al.get('vertical') if al is not None else None,
                    'locked': not (pr is not None and pr.get('locked') in ('0', 'false')), 'hidden': pr is not None and pr.get('hidden') in ('1', 'true'),
                    'border': borders[int(xf.get('borderId') or 0)]})
    dxfs = []
    dx = st.find(M + 'dxfs')
    for d in (dx if dx is not None else []):
        f = d.find(M + 'font')
        fl = d.find(M + 'fill')
        pf = fl.find(M + 'patternFill') if fl is not None else None
        b = f.find(M + 'b') if f is not None else None
        dxfs.append({'fontColor': color_hex(f.find(M + 'color'), theme) if f is not None else None,
                     'bold': b is not None and b.get('val') not in ('0', 'false'),
                     'pattern': pf.get('patternType') if pf is not None else None,
                     'fg': color_hex(pf.find(M + 'fgColor'), theme) if pf is not None else None,
                     'bg': color_hex(pf.find(M + 'bgColor'), theme) if pf is not None else None})
    return {'xfs': xfs, 'dxfs': dxfs}


def parse_sheet(z, path, sst):
    """Return dict with meta and cells (ref -> {v,f,t})"""
    data = z.read(path)
    root = ET.fromstring(data)
    meta = {}
    dim = root.find(M + 'dimension')
    meta['dimension'] = dim.get('ref') if dim is not None else None
    sp = root.find(M + 'sheetProtection')
    meta['protected'] = sp is not None
    sv = root.find(M + 'sheetViews')
    meta['tabColor'] = None
    pr = root.find(M + 'sheetPr')
    if pr is not None:
        tc = pr.find(M + 'tabColor')
        if tc is not None:
            meta['tabColor'] = dict(tc.attrib)
        meta['codeName'] = pr.get('codeName')
    # cols
    cols = []
    c = root.find(M + 'cols')
    if c is not None:
        for col in c:
            cols.append({'min': int(col.get('min')), 'max': int(col.get('max')), 'hidden': col.get('hidden') in ('1', 'true'),
                         'width': col.get('width'), 'outline': col.get('outlineLevel'), 'style': col.get('style')})
    meta['cols'] = cols
    # merges
    merges = []
    mc = root.find(M + 'mergeCells')
    if mc is not None:
        merges = [m.get('ref') for m in mc]
    meta['merges'] = merges
    # DV
    dvs = []
    d = root.find(M + 'dataValidations')
    if d is not None:
        for dv in d:
            f1 = dv.find(M + 'formula1')
            f2 = dv.find(M + 'formula2')
            dvs.append({'type': dv.get('type'), 'op': dv.get('operator'), 'allowBlank': dv.get('allowBlank'), 'sqref': dv.get('sqref'),
                        'f1': f1.text if f1 is not None else None, 'f2': f2.text if f2 is not None else None,
                        'showError': dv.get('showErrorMessage'), 'errStyle': dv.get('errorStyle'), 'errTitle': dv.get('errorTitle'),
                        'err': dv.get('error'), 'promptTitle': dv.get('promptTitle'), 'prompt': dv.get('prompt'), 'x14': False})
    # x14 DV in extLst
    ext = root.find(M + 'extLst')
    cf14 = 0
    if ext is not None:
        for e in ext:
            for dvs14 in e.iter(X14 + 'dataValidations'):
                for dv in dvs14:
                    f1 = dv.find(X14 + 'formula1')
                    f2 = dv.find(X14 + 'formula2')
                    ff1 = f1.find(XM + 'f').text if f1 is not None and f1.find(XM + 'f') is not None else None
                    ff2 = f2.find(XM + 'f').text if f2 is not None and f2.find(XM + 'f') is not None else None
                    sq = dv.find(XM + 'sqref')
                    dvs.append({'type': dv.get('type'), 'op': dv.get('operator'), 'allowBlank': dv.get('allowBlank'),
                                'sqref': sq.text if sq is not None else None, 'f1': ff1, 'f2': ff2,
                                'showError': dv.get('showErrorMessage'), 'errStyle': dv.get('errorStyle'), 'errTitle': dv.get('errorTitle'),
                                'err': dv.get('error'), 'promptTitle': dv.get('promptTitle'), 'prompt': dv.get('prompt'), 'x14': True})
            for cf in e.iter(X14 + 'conditionalFormatting'):
                cf14 += 1
    meta['dataValidations'] = dvs
    cfs = root.findall(M + 'conditionalFormatting')
    meta['condFmtCount'] = len(cfs) + cf14
    meta['condFmt'] = []
    for cf in cfs[:200]:
        rules = []
        for r in cf.findall(M + 'cfRule'):
            fs = [ff.text for ff in r.findall(M + 'formula')]
            rules.append({'type': r.get('type'), 'op': r.get('operator'), 'f': fs, 'dxfId': r.get('dxfId'), 'priority': r.get('priority')})
        meta['condFmt'].append({'sqref': cf.get('sqref'), 'rules': rules})
    # expression-правила (CF за формулою, напр. N($B5="Inactive")) - ВСІ, без обрізання до 200 блоків: аркуш 10 має 489 блоків,
    # а `condFmt` вище лишається як був (його споживачі - cellIs/conditionalFormats - не повинні змінитися).
    meta['condFmtExpr'] = []
    for cf in cfs:
        for r in cf.findall(M + 'cfRule'):
            if r.get('type') == 'expression':
                meta['condFmtExpr'].append({'sqref': cf.get('sqref'), 'f': [ff.text for ff in r.findall(M + 'formula')],
                                            'dxfId': r.get('dxfId'), 'priority': r.get('priority')})
    tp = root.find(M + 'tableParts')
    meta['tableParts'] = len(tp) if tp is not None else 0
    meta['hasDrawing'] = root.find(M + 'drawing') is not None
    meta['hasLegacyDrawing'] = root.find(M + 'legacyDrawing') is not None
    # rows and cells
    rows_meta = {}
    cells = {}
    cell_style = {}
    shared_masters = {}
    sd = root.find(M + 'sheetData')
    for row in sd:
        rn = int(row.get('r'))
        rm = {}
        if row.get('hidden') in ('1', 'true'):
            rm['hidden'] = True
        if row.get('outlineLevel'):
            rm['outline'] = row.get('outlineLevel')
        if rm:
            rows_meta[rn] = rm
        for c in row:
            ref = c.get('r')
            if c.get('s') not in (None, '0'):
                cell_style[ref] = int(c.get('s'))
            t = c.get('t')
            v = c.find(M + 'v')
            f = c.find(M + 'f')
            isel = c.find(M + 'is')
            val = None
            if t == 's' and v is not None:
                val = sst[int(v.text)]
            elif t == 'inlineStr' and isel is not None:
                val = ''.join(x.text or '' for x in isel.iter(M + 't'))
            elif v is not None:
                val = v.text
            fx = None
            if f is not None:
                ft = f.get('t')
                if ft == 'shared':
                    si = f.get('si')
                    if f.text:
                        shared_masters[si] = (ref, f.text)
                        fx = f.text
                    else:
                        if si in shared_masters:
                            mref, mtext = shared_masters[si]
                            try:
                                fx = Translator('=' + mtext, origin=mref).translate_formula(ref)[1:]
                            except Exception:
                                fx = mtext
                else:
                    fx = f.text
                    if ft == 'array':
                        fx = '{' + (fx or '') + '}'
            if val is None and fx is None:
                continue
            cells[ref] = {'v': val, 'f': fx, 't': t}
    meta['rowsMeta'] = rows_meta
    meta['cellStyle'] = cell_style
    return meta, cells


def load(path):
    z = zipfile.ZipFile(path)
    sst = load_shared(z)
    sheets, names = load_workbook_meta(z)
    out = {'sheets': [], 'names': names, 'styles': parse_styles(z)}
    for s in sheets:
        meta, cells = parse_sheet(z, s['path'], sst)
        s.update(meta)
        s['cells'] = cells
        out['sheets'].append(s)
        print(s['name'], s['state'], s['dimension'], len(cells), file=sys.stderr)
    return out


if __name__ == '__main__':
    d = load(sys.argv[1])
    with open(sys.argv[2], 'wb') as f:
        pickle.dump(d, f)
