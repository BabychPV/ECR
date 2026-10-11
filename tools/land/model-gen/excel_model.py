"""Розбір xlsm (FE/BE) у машиночитну модель: аркуші, блоки(таблиці), колонки, DV, CF, формули, довідники."""
import pickle, sys, re, json, collections, hashlib
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

SRC = sys.argv[1] if len(sys.argv) > 1 else 'fe'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'model_%s.json' % SRC
d = pickle.load(open(SRC + '.pkl', 'rb'))
cr = re.compile(r'^([A-Z]+)(\d+)$')
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
MONTHS_RU = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь', 'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']
TITLE_RE = re.compile(r'^\s*\d{1,2}[a-d]?(\.\d+)+\s')


def clean(v):
    return None if v is None else v.replace('_x000D_', '').replace('\r', '')


def one(v):
    return re.sub(r'\s+', ' ', clean(v) or '').strip()


sheetsByName = {s['name']: s for s in d['sheets']}
names = {n['name']: n['text'] for n in d['names'] if n['localSheetId'] is None and n['text']}


def cellv(s, r, c):
    x = s['cells'].get('%s%d' % (get_column_letter(c), r))
    return x['v'] if x else None


def resolve_range(ref):
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", (ref or '').strip())
    if not m or m.group(1) not in sheetsByName:
        return None
    sh, c1, r1, c2, r2 = m.groups()
    s = sheetsByName[sh]
    c2 = c2 or c1; r2 = r2 or r1
    out = []
    for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
        for r in range(int(r1), int(r2) + 1):
            v = cellv(s, r, c)
            if v is not None and one(v) != '':
                out.append(clean(v))
    return out


dictionaries = {}


def register_dict(name, rng, vals):
    key = name or rng
    if key not in dictionaries:
        dictionaries[key] = {'name': name, 'range': rng, 'count': len(vals), 'values': vals}
    return key


def resolve_list(f1):
    if not f1:
        return None, None, None
    f = f1.strip()
    if f in names:
        vals = resolve_range(names[f])
        return (f, names[f], vals) if vals is not None else (f, names[f], None)
    if '!' in f:
        return None, f, resolve_range(f)
    if re.match(r'^\$?[A-Z]+\$?\d+(:\$?[A-Z]+\$?\d+)?$', f):
        return None, f, None
    if f.startswith('"'):
        return None, None, [x.strip() for x in f.strip('"').split(',')]
    return None, None, None


def merge_map(s):
    mm = {}
    for m in s['merges']:
        try:
            c1, r1, c2, r2 = range_boundaries(m)
        except Exception:
            continue
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    mm[(r, c)] = (r1, c1)
    return mm


def norm_formula(f, row):
    if f is None:
        return None
    def rep(m):
        col, dollar, num = m.group(1), m.group(2), int(m.group(3))
        if dollar == '$':
            return m.group(0)
        delta = num - row
        return '%s[r%s]' % (col, ('%+d' % delta) if delta else '')
    return re.sub(r'(?<![A-Za-z_\.])(\$?[A-Z]{1,3})(\$?)(\d+)(?![\d\(])', rep, f)


UNIT_PATTERNS = [
    (r'\(tonn|\(тонн|tonne\)|tonnes\)|\(tonne|\(tons?\)', 't'), (r'tonn/h|tonn/hour|тонн/час', 't/h'),
    (r'\(hours?\)|\(hour|\(час|hours\)', 'h'), (r'\(days?\)|\(дни\)|\(day|\(дней', 'd'),
    (r'sm3|ст\.\s*м3|ст\.м3|ст\. м3', 'Sm3'), (r'\(m3\)|\(м3\)|\(m3|thous m3|тыс\. м3|тыс\.м3', 'm3'),
    (r'kg/h|кг/ч', 'kg/h'), (r'\bkw\b|кВт', 'kW'), (r'g/s|g/sec|г/с', 'g/s'), (r'mg/l|мг/л', 'mg/l'),
    (r'\(rpm|мин-', 'rpm'), (r'\(l/hr|л/ч', 'l/h'), (r'\(kg\)|\(кг\)', 'kg'), (r'\(ea\)|, ea\b|шт\b', 'pcs'),
]


def unit_of(*texts):
    t = ' '.join(x or '' for x in texts).lower()
    for pat, u in UNIT_PATTERNS:
        if re.search(pat, t):
            return u
    return None


def en_ru(text):
    """Розділяє 'EN\\nRU' / 'EN - RU' / 'EN / RU'."""
    t = clean(text) or ''
    if '\n' in t:
        parts = [p.strip() for p in t.split('\n') if p.strip()]
        en = [p for p in parts if not re.search('[А-Яа-яЁё]', p)]
        ru = [p for p in parts if re.search('[А-Яа-яЁё]', p)]
        return ' '.join(en).strip(), ' '.join(ru).strip()
    m = re.match(r'^(.*?)\s+[-–/]\s+(.*[А-Яа-яЁё].*)$', t)
    if m and not re.search('[А-Яа-яЁё]', m.group(1)):
        return m.group(1).strip(), m.group(2).strip()
    if re.search('[А-Яа-яЁё]', t) and not re.search('[A-Za-z]{3}', t):
        return '', t.strip()
    return t.strip(), ''


def find_blocks(s, rows):
    name = s['name']
    maxr = max(rows) if rows else 0
    blocks = []
    if name.startswith('7. Water'):
        return [{'en': 5, 'ru': None, 'top': 1, 'first': 6, 'last': max(r for r in rows if r <= 121), 'title': 'Water Report'}]
    if name.startswith('7.0'):
        return [{'en': 1, 'ru': None, 'top': 1, 'first': 2, 'last': maxr, 'title': 'Water consolidation'}]
    if name.startswith('8. Waste'):
        return [{'en': 1, 'ru': None, 'top': 1, 'first': 3, 'last': maxr, 'title': 'Waste Report'}]
    if name.startswith('9b'):
        return [{'en': 1, 'ru': 2, 'top': 1, 'first': 3, 'last': maxr, 'title': 'Sulfur'}]
    if name.startswith('9c'):
        return [{'en': 2, 'ru': 3, 'top': 1, 'first': 4, 'last': maxr, 'title': 'Fugitive leaks'}]
    if name.startswith('7a') or name.startswith('7b'):
        marks = [r for r in sorted(rows) if one(rows[r].get(1, {}).get('v')) == 'SubstanceId']
        for i, r in enumerate(marks):
            title = ' / '.join(x for x in (one(rows.get(r - 2, {}).get(2, {}).get('v')), one(rows.get(r - 1, {}).get(2, {}).get('v'))) if x)
            blocks.append({'en': r, 'ru': None, 'top': max(1, r - 2), 'first': r + 1, 'title': title or '7a block %d' % (i + 1)})
        for i, b in enumerate(blocks):
            if i + 1 < len(blocks):
                b['last'] = blocks[i + 1]['top'] - 1
            else:
                b['last'] = maxr
        return blocks
    # Status-based
    marks = []
    for r in sorted(rows):
        for c in (2, 3):
            if one(rows[r].get(c, {}).get('v')) == 'Status':
                marks.append(r); break
    for r in marks:
        ru = None
        for rr in range(r + 1, r + 4):
            if one(rows.get(rr, {}).get(2, {}).get('v')).startswith('Работает'):
                ru = rr; break
        blocks.append({'en': r, 'ru': ru or r + 1})
    for i, b in enumerate(blocks):
        prev_end = blocks[i - 1]['ru'] if i > 0 else 0
        title = None
        for rr in range(b['en'] - 1, prev_end, -1):
            t = one(rows.get(rr, {}).get(2, {}).get('v'))
            if TITLE_RE.match(t) and len(t) > 8:
                title = (rr, t); break
        if not title:
            for rr in range(b['en'] - 1, prev_end, -1):
                t = one(rows.get(rr, {}).get(2, {}).get('v'))
                if t and t not in MONTHS and len(t) > 6 and not t.isdigit():
                    title = (rr, t); break
        b['titleRow'] = title[0] if title else None
        b['title'] = title[1] if title else None
        b['top'] = min([x for x in (b['titleRow'], b['en'] - 2) if x] or [b['en']])
        b['top'] = max(b['top'], prev_end + 1, 1)
        b['first'] = b['ru'] + 1
    for i, b in enumerate(blocks):
        nxt = blocks[i + 1]['top'] if i + 1 < len(blocks) else maxr + 1
        b['last'] = nxt - 1
    return blocks


def classify_zone(label):
    l = (label or '').lower()
    if 'tag' in l:
        return 'Tags'
    if 'attribute' in l:
        return 'Attributes'
    if 'measurement' in l:
        return 'Measurements'
    if 'erg' in l:
        return 'ERG'
    if 'limit' in l:
        return 'AnnualLimits'
    if 'remark' in l:
        return 'Remarks'
    if 'promasys' in l or 'eans' in l:
        return 'Service'
    return None


def build_table(s, rows, b, mm):
    en, ru, top = b['en'], b.get('ru'), b['top']
    last_c = 0
    for r in range(top, (ru or en) + 1):
        for c in rows.get(r, {}):
            last_c = max(last_c, c)
    # також колонки з даними у перших рядках даних
    for r in range(b['first'], min(b['first'] + 3, b['last']) + 1):
        for c in rows.get(r, {}):
            last_c = max(last_c, c)
    dvcol = collections.defaultdict(list)
    for dv in s['dataValidations']:
        for rng in (dv['sqref'] or '').split():
            try:
                c1, r1, c2, r2 = range_boundaries(rng)
            except Exception:
                continue
            if r2 < b['first'] or r1 > b['last']:
                continue
            for c in range(c1, min(c2, c1 + 400) + 1):
                dvcol[c].append((dv, r1, r2))
    hidden_cols = set()
    for cd in s['cols']:
        if cd['hidden']:
            for i in range(cd['min'], min(cd['max'], 400) + 1):
                hidden_cols.add(i)

    def hv(r, c):
        src = mm.get((r, c), (r, c))
        return one(cellv(s, src[0], src[1]))

    # зони: у верхніх рядках блоку
    zones = []
    for c in range(1, last_c + 1):
        for r in range(top, en):
            if mm.get((r, c), (r, c)) != (r, c):
                continue
            v = one(cellv(s, r, c))
            if v and v not in MONTHS and v not in MONTHS_RU and v != '*' and len(v) > 3 and classify_zone(v):
                zones.append((c, v))
    zones.sort()

    def zone_of(c):
        z = None
        for zc, zv in zones:
            if zc <= c:
                z = zv
        return z

    data_rows = list(range(b['first'], b['last'] + 1))
    cols = []
    # якорі місяців у верхніх рядках: значення діє до наступного якоря
    anchors = {}
    for r in range(top, en):
        for c in range(1, last_c + 1):
            if mm.get((r, c), (r, c)) != (r, c):
                continue
            g = one(cellv(s, r, c))
            if g in MONTHS:
                anchors.setdefault(c, MONTHS.index(g) + 1)
            elif g in MONTHS_RU:
                anchors.setdefault(c, MONTHS_RU.index(g) + 1)
    month_of_col = {}
    if anchors:
        acs = sorted(anchors)
        for i, ac in enumerate(acs):
            if i + 1 < len(acs):
                end = acs[i + 1] - 1
            else:
                end = ac + ((ac - acs[i - 1]) if i > 0 else 1) - 1
            for c in range(ac, end + 1):
                month_of_col[c] = anchors[ac]
    # логічні колонки: продовження merged-заголовка (en/ru) до складу не входять
    span_of = {}
    cont = set()
    for rr in (en, ru):
        if not rr:
            continue
        for c in range(1, last_c + 1):
            src = mm.get((rr, c))
            if src and src[1] != c and src[0] == rr:
                cont.add(c)
                span_of[src[1]] = span_of.get(src[1], 1) + 1
    for c in range(1, last_c + 1):
        L = get_column_letter(c)
        if c in cont and c not in anchors:
            # продовження merged-заголовка: окрема колонка лише якщо має власний підзаголовок у другому рядку
            own = bool(ru) and one(cellv(s, ru, c)) != '' and mm.get((ru, c), (ru, c)) == (ru, c)
            if not own:
                continue
        h_en = hv(en, c)
        h_ru = hv(ru, c) if ru else ''
        if ru and mm.get((ru, c)) == (en, c):
            h_ru = ''
        glabels = []
        month = month_of_col.get(c)
        for r in range(top, en):
            g = hv(r, c)
            if g:
                glabels.append(g)
            if month is None and g in MONTHS:
                month = MONTHS.index(g) + 1
            elif month is None and g in MONTHS_RU:
                month = MONTHS_RU.index(g) + 1
        base_en, base_ru = en_ru(cellv(s, *(mm.get((en, c), (en, c)))))
        if ru:
            _, r2 = en_ru(cellv(s, *(mm.get((ru, c), (ru, c)))))
            h_ru = h_ru or r2
        if not h_en:
            h_en = base_en
        mname = re.match(r'^(January|February|March|April|May|June|July|August|September|October|November|December)\b', h_en)
        if mname and month is None:
            month = MONTHS.index(mname.group(1)) + 1
        if re.match(r'^[1-4]Q\b', h_en.strip()):
            pass
        nonempty = 0
        formulas = collections.Counter()
        consts = collections.Counter()
        for r in data_rows:
            x = s['cells'].get('%s%d' % (L, r))
            if not x:
                continue
            if x['v'] is None and not x['f']:
                continue
            nonempty += 1
            if x['f']:
                formulas[norm_formula(x['f'], r)] += 1
            elif x['v'] is not None:
                consts[one(x['v'])[:80]] += 1
        dvs = dvcol.get(c, [])
        if not h_en and not h_ru and not nonempty and not dvs and not glabels:
            continue
        col = {'letter': L, 'index': c, 'en': h_en, 'ru': h_ru, 'group': glabels, 'zone': classify_zone(zone_of(c)) if zone_of(c) else None,
               'zoneLabel': zone_of(c), 'hidden': c in hidden_cols, 'month': month, 'filledRows': nonempty,
               'span': span_of.get(c, 1)}
        if formulas:
            f, n = formulas.most_common(1)[0]
            col['formula'] = f
            col['formulaRows'] = sum(formulas.values())
            col['formulaVariants'] = len(formulas)
        if consts:
            col['constants'] = [{'v': k, 'n': n} for k, n in consts.most_common(3)]
            col['distinctConstants'] = len(consts)
        if dvs:
            ents = []
            for dv, r1, r2 in dvs[:3]:
                ent = {'type': dv['type'], 'op': dv['op'], 'f1': dv['f1'], 'f2': dv['f2'], 'allowBlank': dv['allowBlank'],
                       'errorStyle': dv['errStyle'], 'errorTitle': dv['errTitle'], 'error': clean(dv['err']),
                       'promptTitle': dv['promptTitle'], 'prompt': clean(dv['prompt']), 'rows': [r1, r2], 'x14': dv['x14']}
                if dv['type'] == 'list':
                    nm, rg, vals = resolve_list(dv['f1'])
                    ent['listName'] = nm
                    ent['listRange'] = rg
                    ent['listCount'] = len(vals) if vals is not None else None
                    if vals is not None:
                        ent['dictKey'] = register_dict(nm, rg, vals)
                ents.append(ent)
            col['validation'] = ents
        cols.append(col)
    # групові заголовки (merged верхні рядки) як 'headerGroups'
    groups = []
    for m in s['merges']:
        try:
            c1, r1, c2, r2 = range_boundaries(m)
        except Exception:
            continue
        if top <= r1 <= en and r2 < b['first'] and c2 > c1:
            groups.append({'range': m, 'text': one(cellv(s, r1, c1))})
    # merged в даних (кількість)
    dmerged = sum(1 for m in s['merges'] if True and _in_rows(m, b['first'], b['last']))
    # підсумкові рядки
    totals = []
    for r in data_rows:
        for c in (2, 3, 4):
            t = one(rows.get(r, {}).get(c, {}).get('v')).lower()
            if t in ('total', 'итого', 'total:', 'balance', 'input', 'output', 'subtotal') or t.startswith('total'):
                totals.append({'row': r, 'label': t}); break
    return cols, groups, dmerged, totals[:50]


def _in_rows(m, r0, r1):
    try:
        c1, a, c2, b = range_boundaries(m)
    except Exception:
        return False
    return a >= r0 and b <= r1


def classify_sheet(name):
    if name in ('Configuration', 'DropdownList'):
        return 'reference'
    if name.startswith('1.'):
        return 'selector'
    if name.startswith('2.'):
        return 'header'
    return 'report'


model = {'source': SRC, 'sheets': []}
xref = collections.Counter()
for idx, s in enumerate(d['sheets'], 1):
    rows = {}
    for ref, c in s['cells'].items():
        m = cr.match(ref)
        rows.setdefault(int(m.group(2)), {})[column_index_from_string(m.group(1))] = c
        if c['f']:
            for mm_ in re.finditer(r"'([^']+)'!", c['f']):
                xref[(s['name'], mm_.group(1))] += 1
            for mm_ in re.finditer(r"(?<![A-Za-z'])(Configuration|DropdownList)!", c['f']):
                xref[(s['name'], mm_.group(1))] += 1
    mm = merge_map(s)
    hid_rows = sorted(r for r, m in s['rowsMeta'].items() if m.get('hidden'))
    hid_cols = ['%s:%s' % (get_column_letter(cd['min']), get_column_letter(min(cd['max'], 16384))) for cd in s['cols'] if cd['hidden']]
    sh = {'index': idx, 'name': s['name'], 'sheetId': s['sheetId'], 'state': s['state'], 'kind': classify_sheet(s['name']),
          'dimension': s['dimension'], 'protected': s['protected'], 'cellCount': len(s['cells']), 'mergeCount': len(s['merges']),
          'dataValidationCount': len(s['dataValidations']), 'condFormatCount': s['condFmtCount'],
          'hiddenRowCount': len(hid_rows), 'hiddenRows': hid_rows[:100], 'hiddenCols': hid_cols, 'tabColor': s.get('tabColor'),
          'codeName': s.get('codeName'), 'formulaCells': sum(1 for c in s['cells'].values() if c['f']), 'tables': []}
    if sh['kind'] == 'report':
        for ti, b in enumerate(find_blocks(s, rows), 1):
            cols, groups, dm, totals = build_table(s, rows, b, mm)
            sh['tables'].append({'ordinal': ti, 'title': b.get('title'), 'titleRow': b.get('titleRow'), 'headerRow': b['en'], 'headerRowRu': b.get('ru'),
                                 'topRow': b['top'], 'firstDataRow': b['first'], 'lastDataRow': b['last'], 'rowCount': b['last'] - b['first'] + 1,
                                 'headerGroups': groups, 'mergedInData': dm, 'totals': totals, 'columns': cols})
    cfs = []
    for cf in s['condFmt']:
        for r in cf['rules']:
            cfs.append({'sqref': cf['sqref'], 'type': r['type'], 'op': r['op'], 'f': r['f'][:2]})
    sh['conditionalFormats'] = cfs[:400]
    model['sheets'].append(sh)

model['dictionaries'] = dictionaries
model['crossSheetReferences'] = [{'from': a, 'to': b, 'formulas': n} for (a, b), n in sorted(xref.items())]
model['definedNames'] = [{'name': n['name'], 'scope': n['localSheetId'], 'ref': n['text']} for n in d['names']
                         if not n['hidden'] and n['text'] and not n['text'].startswith('#REF') and '[' not in n['text'] and not n['text'].startswith('"')
                         and not re.match(r'^[\d\.]+$', n['text']) and n['text'] not in ('TRUE', 'FALSE')]
json.dump(model, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
tc = tcol = 0
for sh in model['sheets']:
    nc = sum(len(t['columns']) for t in sh['tables'])
    tc += len(sh['tables']); tcol += nc
    print('%-30s %-9s tables=%2d cols=%4d dv=%3d cf=%3d' % (sh['name'], sh['kind'], len(sh['tables']), nc, sh['dataValidationCount'], sh['condFormatCount']))
print('TOTAL tables', tc, 'columns', tcol, 'dictionaries', len(dictionaries))
