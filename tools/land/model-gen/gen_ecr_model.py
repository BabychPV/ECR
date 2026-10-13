"""Excel-модель (model_fe.json + pkl) -> модель побудови шаблону ECR (land-template.model.json)."""
import json, re, pickle, sys, collections, hashlib
from openpyxl.utils import column_index_from_string, get_column_letter

X = json.load(open('model_fe.json', encoding='utf-8'))
PK = pickle.load(open('fe.pkl', 'rb'))
PS = {s['name']: s for s in PK['sheets']}
OUT = sys.argv[1] if len(sys.argv) > 1 else 'land-template.model.json'
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
CYR = re.compile('[А-Яа-яЁё]')


def clean(v):
    return (v or '').replace('_x000D_', '').replace('\r', '')


def one(v):
    return re.sub(r'\s+', ' ', clean(v)).strip()


# Стеля динамічних рядків (MaxDynamicRows) з прод-зразка HSE 2026-10-12 (.sync-local/reports/LAND-PROD-SAMPLE-2026-10-12.md,
# розділ «Прогалини», п. 1). Excel-шаблон має лише 498/184/100 рядків, а типовий документ старої системи — більше:
# T3_00 (обладнання), T9_04, T9_07. Значення = обсяг типового прод-документа із запасом, округлене вгору до 50
# Жорсткої стелі продукту немає: TableDef.SetMaxDynamicRows вимагає лише додатне число, CreateRow/PatchCells порівнюють з
# MaxDynamicRows. Нижню межу (max з рядками Excel і 100) зберігаємо. Інші таблиці зразок не перевищував: «не доведено».
PROD_ROW_LIMITS = {'T3_00': {'maxRows': 1000},
                   'T9_04': {'maxRows': 500},
                   'T9_07': {'maxRows': 250}}


UNIT_PAREN = re.compile(r'\((?:[^)]*\b(?:tonne|tonnes|tonn|hour|hours|hr|Sm3|m3|g/s|g/sec|kg|kg/h|mg/l|days?|rpm|kW|ea|No dimension|No Dimension|L/hr|%|number|tons?)\b[^)]*)\)', re.I)


def slug(s, maxlen=40):
    s = UNIT_PAREN.sub(' ', s or '')
    s = re.sub(r'[()]', ' ', s)
    s = re.sub(r'[^A-Za-z0-9]+', ' ', s).strip()
    parts = s.split()
    out = ''.join(p[:1].upper() + p[1:] for p in parts)
    return out[:maxlen]


def split_enru(v):
    t = one(v)
    m = re.match(r'^(.*?)\s+[-–]\s+(.*)$', t)
    if m and CYR.search(m.group(2)) and not CYR.search(m.group(1)):
        return m.group(1).strip(), m.group(2).strip()
    if CYR.search(t) and not re.search('[A-Za-z]{3}', t):
        return '', t
    return t, ''



# ---------- стилі за ролями Excel (L1, .sync-local/reports/CELL-STYLE-PLAN-2026-10-12.md) ----------
# Кольори беруться ЛИШЕ з xlsm (styles.xml: заливка xf комірки шапки/даних колонки), нічого не вигадано.
# Роль колонки = найчастіший xf її рядків даних; приймається, лише якщо охоплює >= STYLE_MIN_SHARE рядків
# (інакше колонка в Excel неоднорідна, і єдиний стиль колонки в ECR збрехав би - «не доведено», стиль не ставиться).
STYLE_MIN_SHARE = 0.8
DATA_ROLE_BY_FILL = {'CCFFCC': 'AUTO', 'FFFFFF': 'INPUT', 'FFFF00': 'SERVICE', 'FFFF99': 'TAG'}
HEADER_ROLE_BY_FILL = {'CCFFCC': 'HEADER', 'FFCC99': 'HEADER_MONTH'}
FONT_ROLE = {'0000FF': 'BLUE', 'FF0000': 'RED'}   # синій - розрахункові комірки; червоний - перевірки
HALIGN = {'left': 0, 'center': 1, 'right': 2, 'justify': 3}
VALIGN = {'top': 0, 'center': 1, 'bottom': 2}
STYLE_DEFS = {}


def argb_int(rgb):
    v = 0xFF000000 | int(rgb, 16)
    return v - (1 << 32) if v >= (1 << 31) else v


def _style_code(kind, fill, font, bold, halign=None, valign=None, wrap=False):
    if kind == 'header':
        role = HEADER_ROLE_BY_FILL.get(fill) or ('FILL_' + fill if fill else None)
    else:
        role = DATA_ROLE_BY_FILL.get(fill) or ('FILL_' + fill if fill else None)
        if role is None and font == '0000FF':
            role, font = 'CALC', None
    if role is None and not (font or bold):
        return None
    parts = ['LAND', 'HDR' if kind == 'header' and role and not role.startswith('HEADER') else None, role or 'TEXT']
    if font and font in FONT_ROLE:
        parts.append(FONT_ROLE[font])
    if bold:
        parts.append('BOLD')
    code = '_'.join(p for p in parts if p)
    if code not in STYLE_DEFS:
        STYLE_DEFS[code] = {'code': code, 'role': role or 'TEXT', 'backgroundHex': ('#' + fill) if fill else None,
                            'backgroundArgb': argb_int(fill) if fill else None,
                            'foregroundHex': ('#' + font) if font in FONT_ROLE else None,
                            'foregroundArgb': argb_int(font) if font in FONT_ROLE else None,
                            'isBold': bool(bold), 'isItalic': False,
                            'horizontalAlign': None, 'verticalAlign': None, 'wrapText': False, 'columns': 0, 'headers': 0, 'tables': 0, '_al': collections.Counter()}
    if kind == 'header':
        STYLE_DEFS[code]['_al'][(HALIGN.get(halign), VALIGN.get(valign), bool(wrap))] += 1
    return code


def _norm_font(color):
    # лише кольори ролей плану (синій - розрахунок, червоний - перевірка); решта кольорів тексту не переноситься
    return color if color in FONT_ROLE else None


def header_style_code(c):
    h = c.get('hdrStyle')
    if not h or (not h.get('fill') and not h.get('bold')):
        return None
    return _style_code('header', h.get('fill'), _norm_font(h.get('fontColor')), h.get('bold'), h.get('halign'), h.get('valign'), h.get('wrap'))


def style_attrs(c, dtype, hidden):
    """Ролі стилю, ширина й формат числа колонки за xlsm; порожній словник, якщо в Excel нічого не доведено."""
    out = {}
    d = c.get('dataStyle')
    if d and d.get('share', 0) >= STYLE_MIN_SHARE:
        code = _style_code('data', d.get('fill'), _norm_font(d.get('fontColor')), d.get('bold'))
        if code:
            out['styleRole'] = code
            STYLE_DEFS[code]['columns'] += 1
        nf = d.get('numFmt')
        if nf and dtype in ('Decimal', 'Int') and nf not in ('@', 'General') and not nf.startswith('builtin:') and len(nf) <= 50:
            out['displayFormat'] = nf
    hc = header_style_code(c)
    if hc:
        out['headerStyleRole'] = hc
        STYLE_DEFS[hc]['headers'] += 1
    w = c.get('xlWidth')
    if w and not hidden:
        # OOXML <col width> уже включає 5 px відступу: px = width * MDW (7 для Calibri 11); ColumnDef: 40..800 px
        out['widthPx'] = max(40, min(800, int(round(w * 7))))
    return out


# ---------- одиниці ----------
UNIT_MAP = {'t': 't', 'h': 'h', 'd': 'day', 'Sm3': 'Sm3', 'm3': 'm3', 'kg': 'kg', 'g/s': 'g_per_s', 'kg/h': 'kg_per_h', 't/h': 't_per_h',
            'mg/l': 'mg_per_l', 'kW': None, 'rpm': None, 'l/h': None, 'pcs': 'one'}
# factor - до БАЗОВОЇ одиниці розмірності (kg/s для MassFlow, kg/m3, m3), а не до dimensionOf; значення мусять
# збігатися з секцією `-- LAND:units` у 09-seed.sql (G8): kg/h = 1/3600 kg/s, t/h = 1000/3600 kg/s.
# Раніше kg/h і t/h стояли в g/s (x1000 більше) - LandUnitsModelTests стереже збіг із сідом.
NEW_UNITS = [
    {'code': 'kg_per_h', 'symbol': 'kg/h', 'nameEn': 'Kilogram per hour', 'dimensionOf': 'g_per_s', 'factor': '0.000277777777777778'},
    {'code': 't_per_h', 'symbol': 't/h', 'nameEn': 'Tonne per hour', 'dimensionOf': 'g_per_s', 'factor': '0.277777777777777778'},
    {'code': 'mg_per_l', 'symbol': 'mg/l', 'nameEn': 'Milligram per litre', 'dimensionOf': 'kg_per_m3', 'factor': '0.001'},
    {'code': 'thousand_m3', 'symbol': 'thous. m3', 'nameEn': 'Thousand cubic metres', 'dimensionOf': 'm3', 'factor': '1000'},
]
UNIT_PATTERNS = [
    (r'tonn/h|tonn/hour|тонн/час|steaming capacity', 't/h'),
    (r'\(tonn|\(тонн|tonne\)|tonnes\)|\(tonne|\(tons?\)', 't'),
    (r'\(hours?\)|\(hour|\(час|hours\)', 'h'), (r'\(days?\)|\(дни\)|\(day|\(дней|, day', 'd'),
    (r'sm3|ст\.\s*м3|ст\.м3|ст\. м3', 'Sm3'), (r'thous.{0,3}m3|тыс\.\s*м3', 'thousand_m3'), (r'\(m3\)|\(м3\)|\(m3|m3/', 'm3'),
    (r'kg/h|кг/ч', 'kg/h'), (r'\bkw\b|кВт', 'kW'), (r'\(g/s|\(g/sec|g/sec|г/с', 'g/s'), (r'mg/l|мг/л', 'mg/l'),
    (r'\(rpm|мин-', 'rpm'), (r'\(l/hr|л/ч', 'l/h'), (r'\(kg\)|\(кг\)', 'kg'), (r'\(ea\)|, ea\b|шт\b', 'pcs'),
]


def unit_of(*texts):
    t = ' '.join(x or '' for x in texts).lower()
    for pat, u in UNIT_PATTERNS:
        if re.search(pat, t):
            return u
    return None


UNIT_MAP['thousand_m3'] = 'thousand_m3'
UNIT_MAP['one'] = 'one'
# ✎ 2026-10-13 (людина, через координатора): компоненти GTU (301/304/330/337) - мг/м3, O2 - % об'єму. Обидві одиниці
# вже є в базовому сіді (09-seed.sql: mg_per_m3 - MassConcentration, pct_vol - безрозмірна), нових одиниць не додаємо.
UNIT_MAP['mg/m3'] = 'mg_per_m3'
UNIT_MAP['%vol'] = 'pct_vol'

# ---------- довідники ----------
registries = {}
reg_alias = {}  # excel dict key -> registry code


def entry_code(text, used):
    en, _ = split_enru(text)
    base = slug(en or text, 32).upper() or 'E'
    if base[0].isdigit():
        base = 'N' + base
    code = base
    i = 2
    while code in used:
        code = '%s_%d' % (base[:28], i)
        i += 1
    used.add(code)
    return code


def norm_text(v):
    return '\n'.join(re.sub(r'[ \t]+', ' ', ln).strip() for ln in clean(v).split('\n') if ln.strip())


def add_registry(code, name_en, excel, values, fields=None, rows=None, note=None):
    values = [norm_text(v) for v in values]
    code = re.sub(r'[^A-Za-z0-9_]', '_', code)[:64]
    used = set()
    ents = []
    for i, v in enumerate(values):
        en, ru = split_enru(v)
        ents.append({'code': entry_code(v, used), 'name': clean(v).strip(), 'en': en or clean(v).strip(), 'ru': ru})
    r = {'code': code, 'nameEn': name_en, 'excel': excel, 'entries': ents}
    if note:
        r['note'] = note
    registries[code] = r
    return code


SKIP_DICT = re.compile(r'^Contract_')
for key, dct in X['dictionaries'].items():
    if SKIP_DICT.match(key or ''):
        continue
    nm = dct['name'] or re.sub(r'[^A-Za-z0-9]+', '_', dct['range'] or key)
    vals = [v for v in dct['values']]
    if not vals:
        continue
    code = 'LAND_' + re.sub(r'[^A-Za-z0-9]+', '_', nm).upper().strip('_')
    reg_alias[key] = add_registry(code, nm, '%s %s' % (dct['name'] or '', dct['range'] or ''), vals)


# довідники за іменованими діапазонами DropdownList, не задіяні в DV
def named_dd_lists():
    out = {}
    for n in X['definedNames']:
        t = n['ref'] or ''
        if SKIP_DICT.match(n['name']):
            continue
        m = re.match(r"^DropdownList!\$([A-Z]+)\$(\d+)(?::\$([A-Z]+)\$(\d+))?$", t)
        if m and n['name'] not in X['dictionaries']:
            out[n['name']] = t
    return out


dd = PS['DropdownList']


def rng_values(ref):
    m = re.match(r"^(?:DropdownList!)?\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", ref)
    c1, r1, c2, r2 = m.groups()
    c2 = c2 or c1
    r2 = r2 or r1
    vals = []
    for r in range(int(r1), int(r2) + 1):
        x = dd['cells'].get('%s%d' % (c1, r))
        if x and x['v'] not in (None, ''):
            vals.append(clean(x['v']))
    return vals


for name, ref in named_dd_lists().items():
    code = 'LAND_' + re.sub(r'[^A-Za-z0-9]+', '_', name).upper().strip('_')
    if code in registries:
        reg_alias[name] = code
        continue
    vals = rng_values(ref)
    if vals:
        reg_alias[name] = add_registry(code, name, 'DropdownList ' + ref, vals)

# довідники з Configuration
cf = PS['Configuration']


def ccell(ref):
    x = cf['cells'].get(ref)
    return clean(x['v']).strip() if x and x['v'] is not None else ''


def table_rows(cols, r0, r1):
    rows = []
    for r in range(r0, r1 + 1):
        vals = [ccell('%s%d' % (c, r)) for c in cols]
        if any(vals):
            rows.append(vals)
    return rows


def add_table_registry(code, name_en, excel, header, rows, key_idx=0, name_idx=1, ru_idx=None):
    used = set()
    ents = []
    for r in rows:
        key = r[key_idx] or r[name_idx]
        nm = r[name_idx] or key
        ents.append({'code': re.sub(r'[^A-Za-z0-9_]', '_', ('E' + key) if key[:1].isdigit() else key).upper()[:60] or 'E',
                     'name': nm, 'en': nm, 'ru': r[ru_idx] if ru_idx is not None else '', 'fields': dict(zip(header, r))})
    seen = set()
    for e in ents:
        c = e['code']; i = 2
        while c in seen:
            c = '%s_%d' % (e['code'][:55], i); i += 1
        seen.add(c); e['code'] = c
    registries[code] = {'code': code, 'nameEn': name_en, 'excel': excel, 'entries': ents, 'fieldNames': header, 'keyField': header[key_idx]}


add_table_registry('LAND_WATER_REPORT_ENTRY', 'Water report entries (groups, items, balances, notes)', 'Configuration R:U, V:Y, Z:AC, AD:AG',
                   ['ENTRY_ID', 'NAME_EN', 'ENTRY_TYPE', 'NAME_RU'],
                   [[r[0], r[1], r[2], r[3]] for r in table_rows(['S', 'R', 'T', 'U'], 2, 30)] +
                   [[r[0], r[1], r[2], r[3]] for r in table_rows(['V', 'W', 'X', 'Y'], 2, 90)] +
                   [[r[0], r[1], r[2], r[3]] for r in table_rows(['AA', 'Z', 'AB', 'AC'], 2, 10)] +
                   [[r[0], r[1], r[2], r[3]] for r in table_rows(['AE', 'AD', 'AF', 'AG'], 2, 10)], 0, 1, 3)
add_table_registry('LAND_WASTE_REPORT_ENTRY', 'Waste report entries', 'Configuration AJ:AO',
                   ['ENTRY_ID', 'GROUP_ID', 'GROUP_NAME_EN', 'GROUP_NAME_RU', 'NAME_EN', 'NAME_RU'],
                   table_rows(['AJ', 'AK', 'AL', 'AM', 'AN', 'AO'], 2, 60), 0, 4, 5)
add_table_registry('LAND_WASTE_SUBSTANCE', 'Waste substances (type, density, hazard)', 'Configuration AQ:AU',
                   ['SUBSTANCE_CODE', 'NAME_EN', 'NAME_RU', 'DENSITY', 'TYPE_HAZARD'],
                   table_rows(['AQ', 'AR', 'AS', 'AT', 'AU'], 2, 80), 0, 1, 2)
add_table_registry('LAND_PERMIT_WATER', 'Water permits (validity and water body)', 'Configuration AW:BC',
                   ['PERMIT_ID', 'PERMIT_NUMBER', 'START_DATE', 'ACTUAL_END_DATE', 'WATER_BODY_ID', 'NAME_EN', 'EF_TEMPLATE'],
                   table_rows(['AW', 'AX', 'AY', 'AZ', 'BA', 'BB', 'BC'], 2, 60), 1, 5)
wb = []
for c in ['BE', 'BF', 'BG', 'BH', 'BI', 'BJ', 'BK', 'BL']:
    permit = ccell(c + '1')
    for r in range(2, 20):
        n = ccell('%s%d' % (c, r))
        if n:
            wb.append([permit, n])
if wb:
    for i, w in enumerate(wb):
        w.insert(0, 'WB%03d' % (i + 1))
    add_table_registry('LAND_WATER_BODY', 'Water bodies by permit (WBNames)', 'Configuration BE:BL', ['KEY', 'PERMIT_NUMBER', 'NAME_EN'], wb, 0, 2)

CONTRACT_REGISTRIES = ['LAND_AREA', 'LAND_CONTRACTOR', 'LAND_REGION', 'LAND_LOCATION', 'LAND_ONOFF', 'LAND_ACTIVITY', 'LAND_PERMIT']

# ---------- аркуші ----------
SHEET_CODE = {}


def sheet_code(name):
    m = re.match(r'^(\d+)([a-d]?)\.?(\d*)\.?\s*(.*)$', name)
    if not m:
        return 'S_' + slug(name).upper()
    num, letter, sub, rest = m.groups()
    base = 'S%02d' % int(num) + letter.upper()
    if re.match(r'^\d+\.\d', name):
        base += '_' + sub if sub else ''
    return (base + '_' + re.sub(r'[^A-Za-z0-9]+', '_', rest).upper().strip('_'))[:60]


CALC_SHEETS = ('3.', '4.', '5.')
ROW_LABELS = {}

AF = {  # (sheet prefix, table section) -> {normalized header: AF column code}
    '3': {
        'status': 'Land_Status', 'no': 'Land_No', 'subcontractor': 'Land_Subcontractor', 'type of equipment': 'Land_TypeEquipment',
        'description of equipment': 'Land_DescriptionEquipment', 'inventory number of equipment': 'Land_InventoryNumber',
        'techn passport equipment number': 'Land_TechnPassportNumber', 'purpose of equipment': 'Land_Purpose', 'type of fuel': 'Land_TypeFuel',
        'fuel consumption under 100 load l hr from the technical characteristics of the equipment': 'Land_FuelConsumption',
        'power kw or steaming capacity tonn hour': 'Land_Power', 'rotation rpm': 'Land_Rotation', 'write your remarks here': 'Land_Remarks',
        'id code nae project': 'Land_NAE', 'emission source type': 'Land_EmissionSourceType',
        'information on overhauled engines by default not overhaul': 'Land_OverhauledEngines', 'kno2 no dimension': 'Land_EF_KNO2',
        'running hours per month': 'Land_RunningHoursMonthly',
        'liquid fuel consumption': 'Land_Measure_FuelSubstanceConsumptionOrPumpedVolume', 'gas consumption': 'Land_GasConsumption',
    },
    '4': {
        'status': 'Land_Status', 'no': 'Land_No', 'subcontractor': 'Land_Subcontractor', 'equipment type': 'Land_TypeEquipment',
        'type of fuel': 'Land_TypeFuel', 'write your remarks here': 'Land_Remarks', 'id code nae project': 'Land_NAE',
        'total quantity of transport ea': 'Land_Measure_TotQuantTransport', 'fuel consumption kg h': 'Land_Measure_FuelConsumptionKgh',
        'fuel consumption tonne': 'Land_Measure_FuelConsumptionTonn', 'emission source type': 'Land_EmissionSourceType',
    },
    '5': {
        'status': 'Land_Status', 'no': 'Land_No', 'subcontractor': 'Land_Subcontractor', 'inventory number of equipment': 'Land_InventoryNumber',
        'description of equipment': 'Land_DescriptionEquipment', 'tank volume m3': 'Land_TankVolume', 'type of stored material': 'Land_TypeStorMaterial',
        'write your remarks here': 'Land_Remarks', 'id code nae project': 'Land_NAE', 'emission source type': 'Land_EmissionSourceType',
        'quantity of liquid added to tank tonne': 'Land_Measure_QuantLiquid', 'operating time hours': 'Land_Measure_OperTime',
        'operating time days': 'Land_Measure_OperTimeDays',
        'type of petrol station': 'Land_TypePetrStation', 'type of pumped liquid': 'Land_TypePumpedLiquid',
        'quantity of fuel added to petrol station fuel truck tank tonne': 'Land_Measure_QuantFuel',
        'duration of storage of purified water in lagoons day': 'Land_Measure_DurationStorPurWater',
        'type of equipment': 'Land_TypeEquipment', 'type of fuel': 'Land_TypeFuel', 'type of burning fuel': 'Land_TypeBurnFuel',
        'number of ignitions': 'Land_Measure_IgnitionsNumber',
    },
}


PKG = json.load(open('af-land-package.json', encoding='utf-8'))
CALC = {'T3_00': ['ECW_C05_01', 'ECW_C05_02', 'ECW_C05_03', 'ECW_C05_04', 'ECW_C05_05'], 'T4_01': ['ECW_C06_01'],
        'T5_08': ['ECW_C07_07', 'ECW_C07_07_01'], 'T5_09': ['ECW_C07_08'], 'T5_22': ['ECW_C07_16'], 'T5_27': ['ECW_C07_21']}
# Виходи методологій C08/C09/C11/C12 -> таблиці без колонок результатів (LP-08). Файл land-result-outputs.json складає
# extract_result_outputs.py (структура пакета методологій AF x таблиці правил прода, без значень); див. docs/.../land-template/README.md.
RESULT_TABLES = {}
RESULT_TABLES_USED = set()
RESULT_OUTPUTS_PENDING = {}   # таблиці зі зіставлення, яких немає в моделі (чекають LP-09)
try:
    _ro = json.load(open('land-result-outputs.json', encoding='utf-8'))
except FileNotFoundError:
    _ro = {'methodologies': {}}
_ro_all = dict(_ro['methodologies'])
# supplementalMethodologies: виходи, яких немає в пакеті AF XML (ECW_C09_02_01 -> T7b_00); записані зі стенда (структура)
_ro_all.update(_ro.get('supplementalMethodologies', {}))
for _m, _d in sorted(_ro_all.items()):
    for _t in _d['tables']:
        _e = RESULT_TABLES.setdefault(_t, {'methodologies': [], 'outputs': [], 'inputs': []})
        _e['methodologies'].append(_m)
        _e['outputs'] = sorted(set(_e['outputs']) | set(_d['outputs']))
        _e['inputs'] = sorted(set(_e['inputs']) | set(_d.get('inputs', [])))
TEXT_ARGS = {'T3_00': ['Land_OverhauledEngines', 'Land_Rotation'], 'T5_09': ['Land_TypePetrStation'], 'T5_27': ['Land_TypeBurnFuel']}
EXTRA = {
    'T3_00': {'CalculationMonth': ('Formula', '[Period].Sequence', 'one', 'Calculation month'),
              'Density': ('Decimal', None, 'one', 'Gas density (HSE400; manual input here)'),
              'FuelLowerHeating': ('Decimal', None, 'one', 'Lower heating value (HSE400; manual input here)'),
              'H2Smass': ('Decimal', None, 'one', 'H2S mass (HSE400; manual input here)'),
              'MercaptanSulphurContent': ('Decimal', None, 'one', 'Mercaptan sulphur content (HSE400; manual input here)')},
    'T5_08': {'DaysInCalcMonth': ('Formula', '[Period].Days', 'one', 'Days in calculation month'),
              'DaysInCalcYear': ('Formula', '([Period].Year % 4 = 0) ? 366 : 365', 'one', 'Days in calculation year'),
              'Winter': ('Formula', '([Period].Sequence < 3 OR [Period].Sequence > 8) ? 1 : 0', 'one', 'Winter flag')},
    'T5_09': {'Winter': ('Formula', '([Period].Sequence < 3 OR [Period].Sequence > 8) ? 1 : 0', 'one', 'Winter flag')},
    'T5_22': {'CalculationMonth': ('Formula', '[Period].Sequence', 'one', 'Calculation month')},
    'T5_27': {'Con_Region': ('String', None, None, 'Contract region (header Region; manual input here)')},
}


NUMERIC_AF = {'Land_EF_KNO2', 'Land_Power', 'Land_FuelConsumption', 'Land_TankVolume', 'Land_ERG_ExpFactor', 'Land_ERG_StorEmissions',
              'Land_ERG_MaximumVolume', 'Land_ERG_Vsl', 'Land_ERG_DurationAVGdischarge'}


def norm_hdr(h):
    t = clean(h).lower().replace('ч', 'ch')
    t = re.sub(r'[^a-z0-9%]+', ' ', t).strip()
    return t


def af_lookup(prefix, en, ru):
    d = AF.get(prefix, {})
    n = norm_hdr(en)
    if n in d:
        return d[n]
    # префіксні збіги (довгі заголовки)
    for k, v in d.items():
        if len(k) > 12 and (n.startswith(k) or k.startswith(n)) and len(n) > 8:
            return v
        if len(k) > 28 and len(n) > 28 and k[:28] == n[:28]:
            return v
    if prefix == '5':
        if n.startswith('kp max'):
            return 'Land_ERG_ExpFactor'
        if n.startswith('gxp'):
            return 'Land_ERG_StorEmissions'
        if n.startswith('vch max') or n.startswith('v max'):
            return 'Land_ERG_MaximumVolume'
        if n.startswith('average duration of petroleum'):
            return 'Land_ERG_DurationAVGdischarge'
        if n.startswith('vsl') or n.startswith('v sl') or n.startswith('v the volume of oil product'):
            return 'Land_ERG_Vsl'
    return None


def sec_of(title):
    # Номер секції з заголовка таблиці: «9a.16 ...», а також з літерою після номера: «9a.17а» (кирилична а), «9a.21a» (латинська a).
    # Раніше вираз вимагав межі слова одразу після цифр, тож такі заголовки не розпізнавались, а код таблиці брався з
    # порядкового номера блоку (T9a_03/T9a_06/T9a_07 замість T9a_17a/T9a_21a/T9a_21b) - LP-09.
    m = re.match(r'^\s*(\d{1,2}[a-d]?(?:\.\d+)+)([a-d\u0430])?(?!\w)', title or '')
    if not m:
        return None
    return m.group(1) + (m.group(2).replace('\u0430', 'a') if m.group(2) else '')


def table_code(sh_idx, sheet_name, t, k, nblocks):
    sec = sec_of(t.get('title'))
    if sec:
        return 'T' + sec.replace('.', '_')
    pfx = re.match(r'^(\d+[a-d]?)', sheet_name).group(1)
    return 'T%s_%02d' % (pfx, k) if nblocks > 1 else 'T%s_00' % pfx


WATER_OV = {'A': ('No', 'Int'), 'B': ('ReportRowId', 'String'), 'C': ('ReportRowType', 'String'), 'D': ('NameEn', 'String'), 'E': ('NameRu', 'String'),
            'G': ('WaterBodyName', 'String'), 'H': ('OpeningBalance', 'Decimal'), 'V': ('SiteOffshore', 'Int'), 'W': ('SiteAtyrauOnshore', 'Int'),
            'X': ('SiteVessels', 'Int'), 'Y': ('SiteMangistauOnshore', 'Int'), 'Z': ('WaterBodyId', 'String'), 'DF': ('PiTag', 'String'), 'DG': ('PiUnit', 'String')}
M_CODES = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def special(name, c, en, ru):
    L = c['letter']
    mo = c.get('month')
    if name.startswith('7. '):
        if L == 'F':
            return {'code': 'PermitNumber', 'type': 'Lookup', 'registry': 'LAND_PERMITNUMBERS'}
        if L in WATER_OV:
            code, tp = WATER_OV[L]
            d = {'code': code, 'type': tp}
            if L in ('A', 'B', 'C', 'D', 'E'):
                d.update({'hidden': L != 'A', 'readOnly': True, 'represented': 'row definition (key/label)' if L != 'A' else None})
            if L == 'H':
                d['unit'] = 'm3'
            return d
        if mo:
            return {'type': 'Decimal', 'unit': 'm3'}
        return {'type': 'String', 'hidden': True} if not c.get('service') else None
    if name.startswith('7.0'):
        if L in ('A',):
            return {'code': 'No', 'type': 'Int', 'readOnly': True}
        if L == 'B':
            return {'code': 'NameEn', 'type': 'String', 'readOnly': True, 'represented': 'row label'}
        if L == 'C':
            return {'code': 'NameRu', 'type': 'String', 'readOnly': True, 'represented': 'row label'}
        return {'type': 'Decimal', 'unit': 'thousand_m3', 'readOnly': True}
    if name.startswith(('7a', '7b')):
        if L == 'A':
            return {'code': 'SubstanceId', 'type': 'String', 'readOnly': True}
        if L == 'B':
            return {'code': 'No', 'type': 'Int', 'readOnly': True}
        if L == 'C':
            return {'code': 'ComponentEn', 'type': 'String', 'readOnly': True, 'represented': 'row label'}
        if L == 'D':
            return {'code': 'ComponentRu', 'type': 'String', 'readOnly': True, 'represented': 'row label'}
        if L == 'E':
            return {'code': 'AnnualLimit', 'type': 'Decimal', 'unit': 'mg/l'}
        if mo:
            return {'type': 'Decimal', 'unit': 'mg/l'}
        return {'type': 'String', 'hidden': True} if not c.get('service') else None
    if name.startswith('8.'):
        if L == 'A':
            return {'code': 'GroupName', 'type': 'String', 'readOnly': True}
        if L == 'B':
            return {'code': 'ItemName', 'type': 'String', 'readOnly': True, 'represented': 'row label'}
        ci = c['index']
        if 3 <= ci <= 40:
            first = (en or '').split('\n')[0]
            return {'code': 'W_' + (slug(first, 30) or ('C%d' % ci)), 'type': 'Decimal', 'unit': 't'}
        if 41 <= ci <= 44:
            return {'type': 'Int'}
        return {'type': 'String', 'hidden': True} if not c.get('service') else None
    if name.startswith('9b'):
        if c['index'] in (2, 3, 4):
            return {'code': {2: 'Group', 3: 'PadTag', 4: 'PadName'}[c['index']], 'type': 'String', 'readOnly': True}
        if mo or L == 'Q' or 5 <= c['index'] <= 17:
            return {'type': 'Decimal', 'unit': 't'}
    if name.startswith('9c'):
        ov = {'A': ('SheetRds', 'String'), 'B': ('CalculationType', 'String'), 'C': ('UnitName', 'String'), 'D': ('NaeNumber', 'String'),
              'E': ('NaeSub', 'String'), 'F': ('StreamName', 'String'), 'G': ('StreamType', 'String')}
        if L in ov:
            return {'code': ov[L][0], 'type': ov[L][1]}
        if L in ('H', 'I', 'J', 'K'):
            return {'type': 'Decimal', 'unit': 'kg/h'}
        if L in ('L', 'M', 'N', 'O'):
            return {'type': 'Int'}
        if mo:
            return {'type': 'Decimal', 'unit': 'd'}
        if c['index'] >= 28:
            return {'type': 'Decimal'}
    return None


def op_expr(code, v):
    # v: validation entry
    op = v.get('op') or 'between'
    f1, f2 = v.get('f1'), v.get('f2')
    ref = 'SUM([%s])' % code
    try:
        if op == 'lessThanOrEqual': return '%s <= %s' % (ref, f1)
        if op == 'lessThan': return '%s < %s' % (ref, f1)
        if op == 'greaterThanOrEqual': return '%s >= %s' % (ref, f1)
        if op == 'greaterThan': return '%s > %s' % (ref, f1)
        if op == 'equal': return '%s = %s' % (ref, f1)
        if op == 'notEqual': return '%s <> %s' % (ref, f1)
        if op == 'between': return '(%s >= %s AND %s <= %s)' % (ref, f1, ref, f2)
        if op == 'notBetween': return '(%s < %s OR %s > %s)' % (ref, f1, ref, f2)
    except Exception:
        return None
    return None


def op_msg(v):
    op = v.get('op') or 'between'
    f1, f2 = v.get('f1'), v.get('f2')
    words = {'lessThanOrEqual': 'must be <= %s', 'lessThan': 'must be < %s', 'greaterThanOrEqual': 'must be >= %s', 'greaterThan': 'must be > %s',
             'equal': 'must be = %s', 'notEqual': 'must not be %s'}
    if op in words:
        return 'Value ' + words[op] % f1
    if op == 'between':
        return 'Value must be between %s and %s' % (f1, f2)
    return 'Value is out of the allowed range'


def is_num(s):
    try:
        float(s)
        return True
    except Exception:
        return False


out = {'formatVersion': 1, 'sourceFile': 'RDS_xxx_Template_Land_2026_v1.0.4.0.xlsm', 'templateCode': 'LAND_UI',
       'templateNameEn': 'RDS Land 2026 v1.0.4.0 (UI acceptance base)', 'versionNumber': '1.0.4.0',
       'units': {'new': NEW_UNITS, 'map': UNIT_MAP}, 'registries': [], 'contractRegistries': CONTRACT_REGISTRIES, 'sheets': [],
       'excelSheets': [], 'stats': {}}
stat_cols = 0
stat_tables = 0
hidden_default_sheets = []

for sh in X['sheets']:
    es = {'index': sh['index'], 'name': sh['name'], 'kind': sh['kind'], 'ecrSheet': None, 'representedAs': None}
    if sh['kind'] == 'reference':
        es['representedAs'] = 'registries'
    elif sh['kind'] == 'selector':
        es['representedAs'] = 'document sheet selection (sheetDefIds at document creation) + SheetDef.IsMandatory/IsVisible'
    elif sh['kind'] == 'header':
        es['representedAs'] = 'document header fields (HeaderFieldDef, 13 fields) via tools/land/Apply-ContractHeader.ps1'
    out['excelSheets'].append(es)
    if sh['kind'] != 'report':
        continue
    pfx = re.match(r'^(\d+)', sh['name']).group(1)
    scode = sheet_code(sh['name'])
    es['ecrSheet'] = scode
    is_calc = sh['name'].startswith(CALC_SHEETS)
    s_out = {'code': scode, 'excelName': sh['name'], 'nameEn': sh['name'].split('. ', 1)[-1] if '. ' in sh['name'] else sh['name'],
             'ordinal': sh['index'], 'visible': True, 'mandatory': False, 'calc': is_calc, 'tables': []}
    # RU-назва аркуша — з першого заголовка (мапа вручну не потрібна)
    nblocks = len(sh['tables'])
    for k, t in enumerate(sh['tables'], 1):
        tcode = table_code(sh['index'], sh['name'], t, k, nblocks)
        if any(x['code'] == tcode for x in s_out['tables']):
            # Excel 9a має два блоки з однаковим заголовком «9a.21a» (62 і 50 колонок); прод має AF_..._9a_21a і ..._9a_21b, тож
            # другий блок = «b» (позиція між 9a.21a і 9a.22, ті самі 50 колонок і ті самі методології, що в 9a.22/9a.23).
            if tcode[-1] == 'a' and not any(x['code'] == tcode[:-1] + 'b' for x in s_out['tables']):
                tcode = tcode[:-1] + 'b'
            else:
                tcode = '%s_%d' % (tcode, k)
        title_en, title_ru = split_enru(t.get('title') or sh['name'])
        cols = t['columns']
        has_month = any(c.get('month') for c in cols)
        fixed = (sh['name'].startswith(('7. ', '7.0', '7a', '7b', '8.', '9b', '9c')))
        layout = 'PerPeriodInstance' if is_calc else ('MonthsInColumns' if has_month else 'Static')
        t_out = {'code': tcode, 'nameEn': title_en or tcode, 'nameRu': title_ru, 'layout': layout,
                 'rowMode': 'Fixed' if fixed else 'Dynamic',
                 'maxRows': None if fixed else max(t['rowCount'], 100, PROD_ROW_LIMITS.get(tcode, {}).get('maxRows', 0)),
                 'excel': {'sheet': sh['name'], 'headerRow': t['headerRow'], 'firstDataRow': t['firstDataRow'], 'lastDataRow': t['lastDataRow'],
                           'rowCount': t['rowCount'], 'physicalColumns': sum(c.get('span', 1) for c in cols)},
                 'columns': [], 'rows': [], 'rules': [], 'rowValues': {}}
        used = set()
        by_logical = {}
        ctx_month = None
        # ---- колонки
        for c in cols:
            en = one(c['en']); ru = one(c['ru'])
            grp_label = ''
            if en == '*':
                en = ''
            if ru == '*':
                ru = ''
            if ru and not CYR.search(ru) and ru != en:
                # у другому рядку англійський підзаголовок: перший рядок — груповий підпис
                grp_label = en
                en, ru = ru, ''
            if not (en or ru):
                en = ''
            sub_en = en
            mctx = re.match(r'^(January|February|March|April|May|June|July|August|September|October|November|December)\.\s', en)
            if mctx:
                ctx_month = MONTHS.index(mctx.group(1)) + 1
            if (c.get('zoneLabel') or '').lower().startswith('to be filled in by erg') and c['index'] > 0 and ctx_month and c.get('zone') == 'ERG' \
                    and c['letter'] not in ('BB',):
                pass
            grp_abbr = ''
            if grp_label:
                w = re.findall(r'[A-Za-z]+', grp_label)
                grp_abbr = ''.join(x[0].upper() for x in w[:2]) if len(w) >= 2 else (w[0][:3] if w else '')
                if c.get('zone') in ('ERG', 'AnnualLimits') or len(w) >= 3:
                    pass
                if grp_label.lower().startswith(('tags', 'attributes', 'measurements')):
                    grp_abbr = ''
                else:
                    en = '%s: %s' % (re.sub(r'\s*\(.*$', '', grp_label)[:40].strip(), en)
            col = {'excel': c['letter'], 'span': c.get('span', 1), 'month': c.get('month'), 'zone': c.get('zone'),
                   'hidden': bool(c['hidden']), 'headerEn': en, 'headerRu': ru, 'filledRows': c.get('filledRows', 0)}
            # тип
            v = (c.get('validation') or [None])[0]
            n = norm_hdr(en)
            af = af_lookup(pfx, sub_en, ru) if is_calc and c.get('zone') not in ('Tags', 'Attributes', 'Service') else None
            dtype = 'String'
            reg = None
            unit = None
            if c['letter'] == 'A' and not en:
                col['headerEn'] = 'Excel row id'; en = 'Excel row id'; col['hidden'] = True; col['readOnly'] = True
                base = 'ExcelRowId'
            else:
                base = af or ((grp_abbr + '_' if grp_abbr else '') + slug(sub_en, 36))
                if not af and ctx_month and c.get('month') is None and (c.get('zoneLabel') or '').lower().find('erg') >= 0 \
                        and sh['name'].startswith(('3.', '4.', '5.')) and not re.match(r'^(Quarter|[1-4]\s*Quarter)', sub_en):
                    base = '%s_%s' % (M_CODES[ctx_month - 1], slug(sub_en, 30))
                    col['month'] = ctx_month
            if c.get('zone') in ('Tags', 'Attributes') or n in ('xxxx', 'xxx') or (c.get('constants') and c['constants'][0]['v'] == '*' and c.get('distinctConstants') == 1):
                col['hidden'] = True
                col['readOnly'] = True
                col['service'] = True
                if c.get('constants') and c['constants'][0]['v'] == '*':
                    col['default'] = '*'
            if v and v['type'] == 'list' and v.get('dictKey') and v['dictKey'] in reg_alias:
                dtype = 'Lookup'
                reg = reg_alias[v['dictKey']]
            elif v and v['type'] in ('decimal', 'whole'):
                dtype = 'Decimal' if v['type'] == 'decimal' else 'Int'
                col['dv'] = {'op': v.get('op'), 'f1': v.get('f1'), 'f2': v.get('f2'), 'errorStyle': v.get('errorStyle'), 'error': v.get('error')}
            elif v and v['type'] == 'date':
                dtype = 'Date'
            unit = unit_of(en, ru) if dtype in ('String', 'Decimal', 'Int') else None
            if re.search(r'\bkw\b', en.lower()) and re.search(r'tonn', en.lower()):
                unit = None  # потужність кВт або паропродуктивність т/год — змішана одиниця
            if af in NUMERIC_AF and dtype == 'String':
                dtype = 'Decimal'
                unit = unit or ('one' if af in ('Land_EF_KNO2',) else None)
            if dtype == 'String' and unit and not col.get('service'):
                dtype = 'Decimal'
            if dtype == 'String' and not col.get('service'):
                if n == 'no':
                    dtype = 'Int'
                elif n == 'status' and 'LAND_GENERIC_EQUIPSTATUS' in registries:
                    dtype = 'Lookup'; reg = 'LAND_GENERIC_EQUIPSTATUS'
            if dtype == 'String' and c.get('formula') and c.get('formulaRows', 0) >= max(1, c.get('filledRows', 0) // 2):
                col['excelFormula'] = c['formula']
                col['excelFormulaRows'] = c.get('formulaRows')
            if col.get('hidden') and c.get('zone') in ('Tags', 'Attributes'):
                pass
            spx = special(sh['name'], c, en, ru)
            if spx:
                base = spx.get('code') or base
                dtype = spx.get('type', dtype)
                reg = spx.get('registry') if dtype == 'Lookup' else None
                unit = spx.get('unit') if dtype in ('Decimal', 'Int') else None
                if spx.get('hidden'):
                    col['hidden'] = True
                if spx.get('readOnly'):
                    col['readOnly'] = True
                if spx.get('represented'):
                    col['representedAs'] = spx['represented']
                if dtype not in ('Decimal', 'Int'):
                    col.pop('dv', None)
            # код
            if c.get('month'):
                mmn = re.match(r'^(January|February|March|April|May|June|July|August|September|October|November|December)\b', en)
                if mmn and (sh['name'].startswith(('7. ', '7.0', '7a', '7b', '9b', '9c')) or en == mmn.group(1)):
                    code0 = mmn.group(1)
                else:
                    code0 = '%s_M%02d' % (base or 'Col', c['month'])
            else:
                if en in ('Annual',) or re.fullmatch(r'[1-4]Q\s*', en or ''):
                    code0 = 'Annual' if en == 'Annual' else 'Q' + en.strip()[0]
                else:
                    code0 = base or ('C_' + c['letter'])
            if not re.match(r'^[A-Za-z]', code0 or ''):
                code0 = 'C_' + (code0 or c['letter'])
            code = code0[:60]
            if code in used:
                code = ('%s_%s' % (code0[:56], c['letter']))[:64]
            used.add(code)
            col['code'] = code
            col['dataType'] = dtype
            if dtype == 'Lookup':
                col['registry'] = reg
            if unit:
                col['unit'] = unit
            if dtype in ('Decimal',):
                col['precision'] = 28; col['scale'] = 12
            if c.get('validation') and any(x['allowBlank'] in ('0', 0, False) for x in c['validation']):
                col['required'] = False  # allowBlank=0 у шаблоні лише підказка, обов'язковість не примушується
            col.update(style_attrs(c, dtype, col.get('hidden')))
            t_out['columns'].append(col)
        hdr_roles = collections.Counter(x['headerStyleRole'] for x in t_out['columns']
                                        if x.get('headerStyleRole') and not x.get('month') and not x.get('hidden'))
        if not hdr_roles:
            hdr_roles = collections.Counter(x['headerStyleRole'] for x in t_out['columns'] if x.get('headerStyleRole'))
        if hdr_roles:
            t_out['headerStyleRole'] = hdr_roles.most_common(1)[0][0]
        # ---- розрахункові таблиці: «Ref»-колонки, мостові формули, виходи методологій
        if is_calc and tcode in CALC:
            if tcode == 'T4_01':
                for c_ in t_out['columns']:
                    if c_['code'] == 'Land_TypeFuel':
                        c_['dataType'] = 'Lookup'
                        c_['registry'] = 'LAND_GENERIC_TYPE_OF_FUEL'
                        c_.pop('unit', None)
            by_code = {c['code']: c for c in t_out['columns']}
            for refname in TEXT_ARGS.get(tcode, []):
                if refname in by_code:
                    by_code[refname]['code'] = refname + 'Ref'
                    by_code[refname]['refFor'] = refname
            bases = collections.OrderedDict()
            for c in t_out['columns']:
                m_ = re.match(r'^(.*)_M(\d\d)$', c['code'])
                if m_ and c.get('month'):
                    bases.setdefault(m_.group(1), c)
            ordn = 0
            for base_, src in bases.items():
                parts = ['[Period].Sequence = %d ? [%s_M%02d] : ' % (m, base_, m) for m in range(1, 12)]
                expr = ''.join(parts) + '[%s_M12]' % base_
                t_out['columns'].append({'code': base_, 'dataType': 'Formula', 'headerEn': '%s (current period)' % (src['headerEn'] or base_),
                                         'headerRu': '', 'unit': src.get('unit'), 'month': None, 'hidden': False, 'readOnly': True, 'excel': None,
                                         'ecrAdded': 'bridge: selects the month column by [Period].Sequence', 'expression': expr,
                                         'precision': 28, 'scale': 12})
            for refname in TEXT_ARGS.get(tcode, []):
                t_out['columns'].append({'code': refname, 'dataType': 'Formula', 'headerEn': refname + ' (text of the selected entry)', 'headerRu': '',
                                         'hidden': False, 'readOnly': True, 'excel': None, 'month': None, 'ecrAdded': 'text bridge for AF formulas',
                                         'expression': "REGFIELD([%sRef], 'NAME')" % refname})
            for code_, (typ, expr, unit_, hdr) in EXTRA.get(tcode, {}).items():
                col_ = {'code': code_, 'dataType': typ, 'headerEn': hdr, 'headerRu': '', 'hidden': False, 'readOnly': typ == 'Formula', 'excel': None,
                        'month': None, 'ecrAdded': 'calculation argument (not in the Excel sheet)'}
                if code_ == 'DaysInCalcMonth' and 'Land_Measure_OperTimeDays' in bases:
                    # Паритет зі старою системою (12.10): аргумент 7.7 = відпрацьовані дні поточного місяця
                    # (Land_Measure_OperTimeDays), а не календарні дні періоду; при 0 днів методологія дає 0.
                    expr = ''.join('[Period].Sequence = %d ? [Land_Measure_OperTimeDays_M%02d] : ' % (m, m) for m in range(1, 12)) + '[Land_Measure_OperTimeDays_M12]'
                    col_['ecrAdded'] = 'calculation argument = operating days of the current month (legacy parity)'
                if expr:
                    col_['expression'] = expr
                if unit_:
                    col_['unit'] = unit_
                if typ == 'Decimal':
                    col_['precision'] = 28; col_['scale'] = 12
                if code_ == 'Con_Region':
                    col_['default'] = 'Mangistau - Мангистау'
                t_out['columns'].append(col_)
            outs = []
            for mname in CALC[tcode]:
                mm_ = [m for m in PKG['methodologies'] if m['name'] == mname][0]
                for v in mm_['versions']:
                    for f in v['formulas']:
                        if f['name'].startswith('ECW_EC_') and f['name'] not in outs:
                            outs.append(f['name'])
            for o in sorted(outs):
                t_out['columns'].append({'code': o, 'dataType': 'Calculated', 'headerEn': o, 'headerRu': '', 'unit': 't' if '_tons_' in o else 'g/s',
                                         'hidden': False, 'readOnly': True, 'excel': None, 'month': None, 'precision': 28, 'scale': 12,
                                         'ecrAdded': 'methodology output (CalculationBinding)'})
            t_out['calc'] = {'methodologies': CALC[tcode], 'outputs': sorted(outs)}
        elif tcode in RESULT_TABLES:
            # Колонки результатів методологій C08/C09/C11/C12 (LP-08): та сама форма, що й у CALC-таблиць вище
            # (Calculated, readOnly, 28/12, t або g/s, код колонки = код виходу методології).
            rt = RESULT_TABLES[tcode]
            have = {c['code'] for c in t_out['columns']}
            for o in rt['outputs']:
                if o in have:
                    raise SystemExit('result columns: колонка %s уже є в %s' % (o, tcode))
                t_out['columns'].append({'code': o, 'dataType': 'Calculated', 'headerEn': o, 'headerRu': '',
                                         'unit': 't' if re.match(r'^ECW_EC_tons(_|$)', o) else 'g/s',
                                         'hidden': False, 'readOnly': True, 'excel': None, 'month': None, 'precision': 28, 'scale': 12,
                                         'ecrAdded': 'methodology output (CalculationBinding)'})
            t_out['calc'] = {'methodologies': rt['methodologies'], 'outputs': rt['outputs']}
            RESULT_TABLES_USED.add(tcode)
        # ---- рівномірні формули колонок (єдині, що виражаються в діалекті шаблону)
        if sh['name'].startswith('7.0'):
            qmap = {'Q1': (1, 2, 3), 'Q2': (4, 5, 6), 'Q3': (7, 8, 9), 'Q4': (10, 11, 12), 'Annual': tuple(range(1, 13))}
            for c_ in t_out['columns']:
                if c_['code'] in qmap:
                    c_['dataType'] = 'Formula'
                    c_['readOnly'] = True
                    c_['expression'] = 'SUM(%s)' % ', '.join('[%s]' % MONTHS[m_ - 1] for m_ in qmap[c_['code']])
        if sh['name'].startswith('9b'):
            for c_ in t_out['columns']:
                if c_['code'] == 'TOTAL':
                    c_['dataType'] = 'Formula'
                    c_['readOnly'] = True
                    c_['expression'] = 'SUM(%s)' % ', '.join('[%s]' % x for x in MONTHS)
        # ---- рядки Fixed
        s = PS[sh['name']]
        if fixed:
            def cv(L, r):
                x = s['cells'].get('%s%d' % (L, r))
                return clean(x['v']).strip() if x and x['v'] is not None else ''
            first, last = t['firstDataRow'], t['lastDataRow']
            if sh['name'].startswith('7. '):
                parent = None
                for r in range(first, last + 1):
                    ident, typ, nen, nru = cv('B', r), cv('C', r), cv('D', r), cv('E', r)
                    if not (ident or nen):
                        continue
                    key = 'R' + (re.match(r'^(\d{7})', ident).group(1) if re.match(r'^(\d{7})', ident) else str(r))
                    kind = 'Group' if typ == 'GROUP' else ('Balance' if typ == 'BALANCE' else ('Note' if typ == 'NOTE' else 'Item'))
                    if kind == 'Group':
                        parent = key
                    t_out['rows'].append({'key': key, 'en': nen or ident, 'ru': nru, 'kind': kind, 'parent': parent if kind != 'Group' else None, 'excelRow': r})
            elif sh['name'].startswith('7.0'):
                for r in range(first, last + 1):
                    nen, nru = cv('B', r), cv('C', r)
                    if nen:
                        t_out['rows'].append({'key': 'R%02d' % (r - 1), 'en': nen, 'ru': nru, 'kind': 'Item', 'parent': None, 'excelRow': r})
            elif sh['name'].startswith(('7a', '7b')):
                for r in range(first, last + 1):
                    nen, nru, no = cv('C', r), cv('D', r), cv('B', r)
                    if nen or nru:
                        t_out['rows'].append({'key': 'S%s_%d' % (re.sub(r'[^0-9A-Za-z]', '', no) or r, r), 'en': nen or nru, 'ru': nru, 'kind': 'Item', 'parent': None, 'excelRow': r})
            elif sh['name'].startswith('8.'):
                blk = 33
                mon = None
                inblk = {}
                for r in range(first, last + 1):
                    a, b = cv('A', r), cv('B', r)
                    mm_ = re.match(r'^(January|February|March|April|May|June|July|August|September|October|November|December)\b', a)
                    if mm_:
                        mon = MONTHS.index(mm_.group(1)) + 1
                        mk = 'M%02d' % mon
                        t_out['rows'].append({'key': mk, 'en': a, 'ru': '', 'kind': 'Group', 'parent': None, 'excelRow': r})
                        inblk = {}
                        continue
                    if mon is None or not (a or b):
                        continue
                    mk = 'M%02d' % mon
                    code = re.match(r'^(\d{7})', b)
                    label = b or a
                    lab = (code.group(1) if code else re.sub(r'[^A-Za-z0-9]+', '', label)[:16]) or 'ROW'
                    inblk[lab] = inblk.get(lab, 0) + 1
                    key = '%s_%s%s' % (mk, lab, ('_%d' % inblk[lab]) if inblk[lab] > 1 else '')
                    kind = 'Balance' if label in ('Input', 'Output', 'Balance') or a == 'Balance' else ('Note' if a == 'Remarks' else 'Item')
                    t_out['rows'].append({'key': key, 'en': label, 'ru': '', 'kind': kind, 'parent': mk, 'excelRow': r, 'group': a})
            elif sh['name'].startswith('9b'):
                for r in range(first, last + 1):
                    b, c_, d = cv('B', r), cv('C', r), cv('D', r)
                    if b or c_ or d:
                        t_out['rows'].append({'key': 'R%02d' % r, 'en': (d or c_ or b), 'ru': '', 'kind': 'Item', 'parent': None, 'excelRow': r, 'group': b})
            elif sh['name'].startswith('9c'):
                for r in range(first, last + 1):
                    vals = [cv(x, r) for x in ('A', 'C', 'D', 'E', 'F')]
                    if any(vals):
                        t_out['rows'].append({'key': 'R%03d' % r, 'en': ' / '.join(x for x in vals[1:] if x) or vals[0], 'ru': '', 'kind': 'Item', 'parent': None, 'excelRow': r})
            t_out['maxRows'] = None
            # константні значення у даних Fixed-таблиць (ліміти тощо) — для наповнення документів
            colmap = {c['excel']: c['code'] for c in t_out['columns'] if c.get('excel')}   # без доданих ECR колонок (excel=None)
            rowmap = {rw['excelRow']: rw['key'] for rw in t_out['rows']}
            for L, ccode in colmap.items():
                cd = [c for c in cols if c['letter'] == L][0]
                if cd.get('service') or not cd.get('constants') or cd.get('month'):
                    continue
            for r, rk in rowmap.items():
                vals = {}
                for L, ccode in colmap.items():
                    x = s['cells'].get('%s%d' % (L, r))
                    if x and x['v'] not in (None, '') and not x['f']:
                        cdef = [cc for cc in t_out['columns'] if cc['excel'] == L][0]
                        if cdef.get('service') or cdef['code'] in ('NameEn', 'NameRu'):
                            continue
                        if cdef['dataType'] in ('Decimal', 'Int') and is_num(x['v']):
                            vals[ccode] = x['v']
                if vals:
                    t_out['rowValues'][rk] = vals
        # ---- правила валідації з DV
        seen_rules = set()
        for col in t_out['columns']:
            if col.get('dv') and col['dataType'] in ('Decimal', 'Int'):
                dv = col['dv']
                expr = op_expr(col['code'], dv)
                if expr:
                    rc = ('DV_' + col['code'])[:64]
                    if rc in seen_rules:
                        continue
                    seen_rules.add(rc)
                    sev = 'Error' if (dv.get('errorStyle') in (None, 'stop')) else ('Warning' if dv.get('errorStyle') == 'warning' else 'Info')
                    t_out['rules'].append({'code': rc, 'severity': sev, 'scope': 0, 'expression': expr, 'message': op_msg(dv), 'column': col['code'],
                                           'excelColumn': col['excel']})
        if sh['name'].startswith('8.'):
            def is_text_row(rw):
                return ('Name of the facility' in rw['en']) or rw['en'].startswith('Contract number') or rw['kind'] == 'Note'
            text_rows = [rw for rw in t_out['rows'] if is_text_row(rw)]
            keep = [rw for rw in t_out['rows'] if not is_text_row(rw)]
            keys = {rw['key'] for rw in keep}
            grp = {rw['parent'] for rw in keep if rw['parent']}
            t_out['rows'] = [rw for rw in keep]
            # групові рядки місяців лишаються в обох таблицях
            groups_ = [rw for rw in t_out['rows'] if rw['kind'] == 'Group']
            t2 = json.loads(json.dumps(t_out))
            t2['code'] = 'T8_01'
            t2['nameEn'] = 'Waste report: facility names, contract, remarks (text rows)'
            t2['rows'] = groups_ + text_rows
            for cc in t2['columns']:
                if cc['dataType'] in ('Decimal', 'Int') and cc['code'].startswith('W_'):
                    cc['dataType'] = 'String'
                    cc.pop('unit', None); cc.pop('precision', None); cc.pop('scale', None); cc.pop('dv', None)
            t2['rules'] = []
            t2['rowValues'] = {}
            t_out['nameEn'] = 'Waste report: amounts (tonnes)'
            s_out['tables'].append(t_out)
            s_out['tables'].append(t2)
            stat_cols += len(t_out['columns']) + len(t2['columns'])
            stat_tables += 2
        else:
            s_out['tables'].append(t_out)
            stat_cols += len(t_out['columns'])
            stat_tables += 1
    out['sheets'].append(s_out)

# ---- D1: перехресні формули 7.0 <- 7. Water Report -> формули РІВНЯ КОМІРКИ (PUT …/formulas/cell/{rowKey}/{columnDefId})
# Excel: D..O = (сума конкретних рядків аркуша 7 у тій самій місячній колонці) / 1000. Формула лишається в діалекті
# шаблону: адреси A1 переписуються на [Аркуш].[Таблиця].[RowKey].[Колонка]; діапазонів немає. Усе, що не
# розв'язується (рядок/колонка поза моделлю, інший аркуш), НЕ мовчить — потрапляє в cellFormulasUnresolved.
SUM_FORM = re.compile(r"^\(?(?P<ops>\[[^\]]+\](?:\.\[[^\]]+\]){3}(?:\+\[[^\]]+\](?:\.\[[^\]]+\]){3})*)\)?/1000$")


def _cell_formulas_7_0():
    s70 = next((s for s in out['sheets'] if s['excelName'].startswith('7.0')), None)
    s7 = next((s for s in out['sheets'] if s['excelName'].startswith('7. ')), None)
    if not s70 or not s7:
        return
    t70, t7 = s70['tables'][0], s7['tables'][0]
    src_name = s7['excelName']
    rowkey_by_excel = {rw['excelRow']: rw['key'] for rw in t7['rows']}
    col_by_letter = {c['excel']: c['code'] for c in t7['columns'] if c.get('excel')}
    dst_row_by_excel = {rw['excelRow']: rw['key'] for rw in t70['rows']}
    dst_col_by_letter = {c['excel']: c['code'] for c in t70['columns'] if c.get('excel')}
    ref_re = re.compile(r"'%s'!\$?([A-Z]+)\$?(\d+)" % re.escape(src_name))
    sheet70 = PS[s70['excelName']]['cells']
    res, bad = [], []
    for L in 'DEFGHIJKLMNO':
        mcode = dst_col_by_letter.get(L)
        for r, rkey in sorted(dst_row_by_excel.items()):
            x = sheet70.get('%s%d' % (L, r))
            if not x or not x.get('f'):
                continue
            f = x['f']
            problems = []

            def repl(m):
                col, row = m.group(1), int(m.group(2))
                k, c = rowkey_by_excel.get(row), col_by_letter.get(col)
                if not k or not c:
                    problems.append('%s%d' % (col, row))
                    return m.group(0)
                if c != mcode:
                    problems.append('%s%d (month %s != %s)' % (col, row, c, mcode))
                return '[%s].[%s].[%s].[%s]' % (s7['code'], t7['code'], k, c)

            expr = ref_re.sub(repl, f)
            if re.search(r"'[^']+'!", expr):
                problems.append('reference to another sheet')
            if problems or not mcode:
                bad.append({'cell': '%s%d' % (L, r), 'formula': f, 'why': 'unresolved: ' + ', '.join(problems or ['target column'])})
                continue
            # LP-31: порожня клітинка в Excel = 0, а в ECR `a+b` з відсутнім операндом дає null (ТЗ 05-engines «Семантика
            # null»). Тому `(a+b+c)/1000` і `a/1000` пишемо як `SUM(a,b,c)/1000` / `SUM(a)/1000`: SUM поглинає null,
            # порожня множина = 0 (02b §6.1) — як Excel. Форма, яку не впізнано, НЕ лишається мовчки: unresolved.
            m_sum = SUM_FORM.match(expr)
            if not m_sum:
                bad.append({'cell': '%s%d' % (L, r), 'formula': f, 'why': 'unexpected form (not (a+b+..)/1000): ' + expr})
                continue
            expr = 'SUM(%s)/1000' % ','.join(m_sum.group('ops').split('+'))
            res.append({'row': rkey, 'column': mcode, 'expression': expr, 'excelCell': '%s%d' % (L, r), 'excelFormula': f})
    # Колонка з комірочною формулою мусить бути обчислюваною (ECR-TMPL-4227), інакше PUT …/formulas/cell відхиляється.
    covered = {x['column'] for x in res}
    for c_ in t70['columns']:
        if c_['code'] in covered:
            c_['dataType'] = 'Formula'
            c_['readOnly'] = True
            c_.pop('dv', None)
    t70['cellFormulas'] = res
    t70['cellFormulasUnresolved'] = bad
    print('D1 cellFormulas 7.0:', len(res), 'unresolved:', len(bad))


_cell_formulas_7_0()

# WW-02: cell-формули Water Report / 7a / 7b / Waste Report (групові підсумки, баланси, «Volume», 8001002, Input/Output/Balance).
import cellformulas_report  # noqa: E402
_ww02 = cellformulas_report.apply(out, PS)
print('WW-02 cellFormulas total:', _ww02['total'], 'unresolved:', _ww02['unresolved'], 'skipped(lookups):', _ww02['skipped'])

# ---- відповідність прод-колонок AF і колонки введення, яких немає в Excel (прод-зразок 2026-10-12, прогалини 3-5): див. prod_mapping.py
import prod_mapping  # noqa: E402
out['prodColumnMapping'], _added_cols = prod_mapping.apply(out)
stat_cols += _added_cols
print('prodColumnMapping: tables %d, new columns %d' % (len(out['prodColumnMapping']['tables']), _added_cols))


# ---- мости аргументів методологій C08/C09/C11/C12 для таблиць аркушів 6, 9-11 (LP-10).
# Для аркушів 3-5 (PerPeriodInstance) місячна колонка `X_M01..X_M12` має Formula-міст `X`, що вибирає місяць за
# [Period].Sequence; аргумент методології (імена з AF) знаходить колонку за кодом. Таблиці 6, 9-11 (MonthsInColumns) мостів
# не мали. Тут мости робляться ЛИШЕ там, де відповідність доведена prodColumnMapping (alias прод-колонки AF -> колонка(и)
# шаблону, evidence header/single/position) для входу, який справді використовує методологія таблиці (inputs з
# land-result-outputs.json), та для календарних аргументів (ті самі формули, що в EXTRA таблиць 3-5). Усе інше — у
# argumentBridges.unresolved з причиною: нічого не вгадуємо.
def _arg_bridges():
    global stat_cols
    calendar = {'CalculationMonth': EXTRA['T3_00']['CalculationMonth'], 'DaysInCalcYear': EXTRA['T5_08']['DaysInCalcYear'],
                'Winter': EXTRA['T5_08']['Winter']}
    by_code = {}
    for s_ in out['sheets']:
        for t_ in s_['tables']:
            by_code.setdefault(t_['code'], (s_, t_))
    ptabs = out['prodColumnMapping']['tables']
    bridged, unresolved, added = {}, {}, 0

    def month_expr(cols_):
        return ''.join('[Period].Sequence = %d ? [%s] : ' % (m, cols_[m - 1]) for m in range(1, 12)) + '[%s]' % cols_[11]

    for tcode in sorted(RESULT_TABLES):
        if tcode not in by_code or tcode in CALC:
            continue
        s_, t_ = by_code[tcode]
        if t_['layout'] != 'MonthsInColumns':
            continue
        cols = {c['code']: c for c in t_['columns']}
        info = ptabs.get('%s|%s' % (s_['code'], tcode))
        alias = {a['prod']: a for a in info['aliases']} if info else {}
        done, left = [], {}
        for name in RESULT_TABLES[tcode].get('inputs', []):
            if name in cols:
                continue   # колонка з такою назвою вже є (пряме відповідання)
            new = None
            if name in calendar:
                typ, expr, unit_, hdr = calendar[name]
                new = {'code': name, 'dataType': typ, 'headerEn': hdr, 'unit': unit_, 'expression': expr,
                       'ecrAdded': 'calculation argument (not in the Excel sheet)'}
            elif name == 'DaysInCalcMonth':
                # Правило «DaysInCalcMonth = Land_Measure_OperTimeDays» (7.7) доведено для аркушів 3-5; для 6, 9-11 не доведено
                # (календарні дні [Period].Days чи відпрацьовані дні - рішення власника методології): мост не створюємо.
                left[name] = 'правило 7.7 (відпрацьовані дні чи календарні) доведено лише для аркушів 3-5; для цієї таблиці питання власнику'
            elif name in alias:
                a = alias[name]
                if a.get('monthColumns'):
                    mc = a['monthColumns']
                    ok = len(mc) == 12 and all(c in cols for c in mc) and [cols[c].get('month') for c in mc] == list(range(1, 13))
                    if ok:
                        new = {'code': name, 'dataType': 'Formula', 'headerEn': '%s (argument bridge: selects the month column)' % name,
                               'unit': cols[mc[0]].get('unit'), 'expression': month_expr(mc),
                               'ecrAdded': 'bridge: methodology argument %s = prod column alias (evidence %s)' % (name, a['evidence'])}
                    else:
                        left[name] = 'alias є, але місячні колонки не збігаються з 12 колонками таблиці'
                elif a.get('column') in cols and cols[a['column']]['dataType'] in ('Decimal', 'Int'):
                    src = cols[a['column']]
                    new = {'code': name, 'dataType': 'Formula', 'headerEn': '%s (argument bridge: %s)' % (name, src['code']),
                           'unit': src.get('unit'), 'expression': '[%s]' % src['code'],
                           'ecrAdded': 'bridge: methodology argument %s = prod column alias (evidence %s)' % (name, a['evidence'])}
                else:
                    left[name] = 'alias є, але колонка шаблону не числова (мост Formula неможливий без рішення замовника)'
            else:
                left[name] = 'відповідності прод-колонці в prodColumnMapping немає (не доведено)'
            if new is not None:
                col_ = {'headerRu': '', 'hidden': False, 'readOnly': True, 'excel': None, 'month': None, 'precision': 28, 'scale': 12}
                col_.update({k: v for k, v in new.items() if v is not None})
                t_['columns'].append(col_)
                done.append(name)
                added += 1
        if done:
            bridged[tcode] = sorted(done)
        if left:
            unresolved[tcode] = dict(sorted(left.items()))
    out['argumentBridges'] = {'bridged': dict(sorted(bridged.items())), 'unresolved': dict(sorted(unresolved.items()))}
    stat_cols += added
    print('argumentBridges: tables %d, columns %d, unresolved inputs %d' % (len(bridged), added, sum(len(v) for v in unresolved.values())))


_arg_bridges()

# ---- склад газу HSE400 (LP-10): довідники LAND_HSE400_STREAM / LAND_HSE400_COMPOSITION і мости аргументів Density20, Vddg ...
# Значень складу в моделі немає (їх одноразово завантажує tools/land/Import-LandHse400Composition.ps1); див. hse400_bridges.py.
import hse400_bridges  # noqa: E402
import af_attr_bridges  # noqa: E402
_af_attrs = json.load(open('af-land-table-attributes.json', encoding='utf-8'))['tables']
_cat_rules = json.load(open('af-land-category-rules.json', encoding='utf-8'))['rules']
# джерело тексту потоку для таблиць 9a/10_03/11_03 - за атрибутом AF, який читає старий код (af_attr_bridges.stream_sources)
stat_cols += hse400_bridges.apply(out, registries, af_attr_bridges.stream_sources(out, _af_attrs), _af_attrs)

# ---- застосовність речовин на рядок (рішення людини 13.10): довідник LAND_SUBSTANCE_APPLICABILITY зі словників
# старого C# (land-substance-applicability.json). Без міграції; значень прода немає; див. substance_applicability.py.
import substance_applicability  # noqa: E402
substance_applicability.apply(out, registries, _cat_rules)

# ---- мости за атрибутами AF-таблиць (LP-10, прив'язки C11/C12 -> ECR-CALC-0438/0422): af_attr_bridges.py.
stat_cols += af_attr_bridges.apply(out, _ro_all, _cat_rules, _af_attrs, set(CALC))

# ---- похідні аргументи зі словників C# (Land_Power, Category1 - довідник LAND_NAE_GROUP; Land_MassFraction - мости
# за кодом речовини і фрагмент): nae_group_bridges.py. Після af_attr_bridges: потребує мосту Land_NAE і записів unresolved.
import nae_group_bridges  # noqa: E402
stat_cols += nae_group_bridges.apply(out, registries, _af_attrs)

cfs = []
seen_cf = set()
for sh in X['sheets']:
    for cf in sh.get('conditionalFormats', []):
        if cf['type'] == 'cellIs' and cf['op'] == 'equal' and cf['f'] and cf['f'][0].strip('"') in ('Removed', 'Inactive'):
            val = cf['f'][0].strip('"')
            for code in ('Status', 'Land_Status'):
                if (code, val) not in seen_cf:
                    seen_cf.add((code, val))
                    cfs.append({'columnCode': code, 'operator': 'eq', 'value': val, 'backgroundHex': '#D9D9D9', 'foregroundHex': '#7F7F7F', 'isBold': False})
# ---- масштаб числових колонок: сховище комірки — decimal(34,16) (CellValueConfiguration, CellValueReader.StorageScale = 16).
# Прод-зразок 2026-10-12 (прогалина 10): прод має 16 знаків, а колонки моделі мали (28,12) — ColumnDef.ValidateValue
# відхиляє знак понад Scale, тож значення довелось округляти до 10 знаків. Вхідні Decimal і Formula-колонки (мости/суми, що
# копіюють вхід) отримують (34,16) = межа сховища: без міграції, без втрати знаків. Calculated (виходи методологій) лишаються
# (28,12): їх масштаб задає прив'язка методології (binding.OutputScales), а Scale колонки лише подача.
DECIMAL_PRECISION, DECIMAL_SCALE = 34, 16
for s_ in out['sheets']:
    for t_ in s_['tables']:
        for c_ in t_['columns']:
            if c_.get('dataType') in ('Decimal', 'Formula') and c_.get('precision') is not None:
                c_['precision'] = DECIMAL_PRECISION
                c_['scale'] = DECIMAL_SCALE

# ---- відповідність прод-відходів (AF_Waste_Land_08) -> T8_00/T8_01 шаблону (прод-зразок 2026-10-12, прогалина 2).
# Прод: рядок = (елемент звіту × код відходу), Land_ItemId (= ENTRY_ID довідника LAND_WASTE_REPORT_ENTRY), Land_WasteNo 1..38,
# Land_WasteCode (= SUBSTANCE_CODE довідника LAND_WASTE_SUBSTANCE), значення — Land_Measure_Weight (т). Шаблон: широка таблиця,
# рядок M<mm>_<ENTRY_ID>, колонка W_<відхід>. Колонку відходу визначає НАЗВА: заголовок W-колонки починається з NAME_EN
# відходу (найдовший збіг); мапа мусить бути бієкцією 38 <-> 38, інакше генерація падає (нічого не вгадуємо).
# Незалежне підтвердження: на прод-зразку Land_WasteNo n -> Land_WasteCode збігся з n-ю W-колонкою для всіх 38.
def _waste_mapping():
    sw = next(s for s in out['sheets'] if s['excelName'].startswith('8.'))
    t_amt = next(t for t in sw['tables'] if t['code'] == 'T8_00')
    reg = registries['LAND_WASTE_SUBSTANCE']
    ent = registries['LAND_WASTE_REPORT_ENTRY']
    wcols = [c for c in t_amt['columns'] if c['code'].startswith('W_')]

    def nrm(s):
        return re.sub(r'[^a-z0-9]+', '', (s or '').lower())

    by_col, used = {}, {}
    for c in wcols:
        h = nrm(c['headerEn'])
        hits = [(len(nrm(e['fields']['NAME_EN'])), e) for e in reg['entries']
                if nrm(e['fields']['NAME_EN']) and h.startswith(nrm(e['fields']['NAME_EN']))]
        if not hits:
            raise SystemExit('waste mapping: колонка %s без відходу за назвою' % c['code'])
        best = max(hits, key=lambda x: x[0])
        if sum(1 for n_, _ in hits if n_ == best[0]) != 1:
            raise SystemExit('waste mapping: неоднозначний збіг для %s' % c['code'])
        code = best[1]['fields']['SUBSTANCE_CODE']
        if code in used:
            raise SystemExit('waste mapping: відхід %s двічі (%s, %s)' % (code, used[code], c['code']))
        used[code] = c['code']
        by_col[c['code']] = code
    if len(used) != len(reg['entries']):
        raise SystemExit('waste mapping: %d колонок, %d відходів у довіднику' % (len(used), len(reg['entries'])))
    items = [e['fields']['ENTRY_ID'] for e in ent['entries']]
    groups = collections.OrderedDict()
    for e in ent['entries']:
        groups.setdefault(e['fields']['GROUP_ID'], []).append(e['fields']['ENTRY_ID'])
    # рядки січня як представник місяця (структура місяців однакова)
    amt_items = [rw['key'][4:] for rw in t_amt['rows'] if rw['key'].startswith('M01_') and rw['kind'] == 'Item']
    txt = next(t for t in sw['tables'] if t['code'] == 'T8_01')
    txt_items = [rw['key'][4:] for rw in txt['rows'] if rw['key'].startswith('M01_')]
    return {
        'prodTable': 'AF_Waste_Land_08',
        'template': {'sheet': sw['code'], 'amountTable': 'T8_00', 'textTable': 'T8_01'},
        'rowKey': 'M{month:02d}_{Land_ItemId}',
        'valueProdColumn': 'Land_Measure_Weight', 'valueUnit': 't',
        'wasteCodeProdColumn': 'Land_WasteCode', 'wasteNoProdColumn': 'Land_WasteNo',
        'wasteRegistry': 'LAND_WASTE_SUBSTANCE', 'itemRegistry': 'LAND_WASTE_REPORT_ENTRY',
        'columnByWasteCode': {by_col[c['code']]: c['code'] for c in wcols},
        'wasteNoByCode': {by_col[c['code']]: i for i, c in enumerate(wcols, 1)},
        'itemsInAmountTable': amt_items,
        'itemsInTextTable': txt_items,
        'groupItems': groups,
        'notProven': [
            {'prodColumn': 'Land_Waste_Name_Area', 'why': 'Звіт: значення псевдонімізоване (area-NNN) і не зіставлене з рядками T8_01 '
                                                          '(назва об\'єкта/договір/примітки) — відповідність не доведено, потрібне питання замовнику.'},
            {'prodColumn': 'Land_Value', 'why': 'Текстовий дубль Land_Measure_Weight у зразку (число як текст); у шаблон не переноситься, '
                                                 'джерелом значення береться Land_Measure_Weight.'},
            {'prodColumn': 'Land_No', 'why': 'У зразку лише порядковий номер; ключ рядка шаблону — Land_ItemId.'},
        ],
    }


out['prodMapping'] = {'formatVersion': 1, 'source': 'LAND-PROD-SAMPLE-2026-10-12 (структура; даних прод-БД немає)', 'waste': _waste_mapping()}

out['conditionalFormats'] = []  # ECR підтримує лише числові пороги (eq/ne/gt/...); Excel-правила за текстом «Removed» не переносяться
out['conditionalFormatsNotTransferred'] = cfs

# ---- стилі за ролями (ключ `styles`) і умовне форматування cellIs з прив'язкою до рядків/колонок ECR
# `conditionalFormats` лишається порожнім: правило ECR діє на колонку з цим кодом у ВСІХ таблицях версії й на всі її рядки, а правила
# Excel (баланс <>0 / =0, статус Removed) стоять на конкретних рядках - перенесення на колонку пофарбувало б усі комірки колонки.
# Тому вони лежать окремо (`conditionalFormatsRowScoped`) і застосовуються, коли сервер вміє scope = row (L6 плану).
from openpyxl.utils import range_boundaries as _rb


def _cf_scoped():
    res = []
    for sh in out['sheets']:
        xs = next((x for x in X['sheets'] if x['name'] == sh['excelName']), None)
        if xs is None:
            continue
        for cf in xs.get('cellIsFormats', []):
            op = {'equal': 'eq', 'notEqual': 'ne', 'greaterThan': 'gt', 'lessThan': 'lt', 'greaterThanOrEqual': 'ge', 'lessThanOrEqual': 'le'}.get(cf['op'])
            val = (cf['f'][0] if cf['f'] else '')
            text = val.startswith('"')
            if op is None or (not cf['fill'] and not cf['font']):
                continue
            val = val.strip('"')
            for part in (cf['sqref'] or '').split():
                try:
                    c1, r1, c2, r2 = _rb(part)
                except Exception:
                    continue
                for t in sh['tables']:
                    ex = t['excel']
                    if r2 < ex['firstDataRow'] or r1 > ex['lastDataRow']:
                        continue
                    cols = [c['code'] for c in t['columns'] if c.get('excel') and c1 <= column_index_from_string(c['excel']) <= c2]
                    if not cols:
                        continue
                    keys = [r['key'] for r in t['rows'] if r1 <= r.get('excelRow', 0) <= r2]
                    if t['rowMode'] == 'Fixed' and not keys:
                        continue
                    res.append({'sheet': sh['code'], 'table': t['code'], 'excelRange': part, 'columns': cols, 'rowKeys': keys if t['rowMode'] == 'Fixed' else None,
                                'excelRows': [max(r1, ex['firstDataRow']), min(r2, ex['lastDataRow'])], 'operator': op, 'value': val,
                                'valueKind': 'text' if text else 'number', 'backgroundHex': ('#' + cf['fill']) if cf['fill'] else None,
                                'foregroundHex': ('#' + cf['font']) if cf['font'] else None, 'isBold': bool(cf['bold'])})
    return res


for _s in out['sheets']:
    for _t in _s['tables']:
        if _t.get('headerStyleRole'):
            STYLE_DEFS[_t['headerStyleRole']]['tables'] += 1
for _d in STYLE_DEFS.values():
    # вирівнювання/перенос шапки - найчастіше серед шапок цієї ролі (рахується з xlsm, не припускається)
    if _d['_al']:
        (_d['horizontalAlign'], _d['verticalAlign'], _d['wrapText']), _n = _d['_al'].most_common(1)[0]
    del _d['_al']
out['styles'] = [STYLE_DEFS[k] for k in sorted(STYLE_DEFS)]
out['conditionalFormatsRowScoped'] = _cf_scoped()


# ---- умовне форматування за СТАТУСОМ РЯДКА: expression-правила Excel N($B5="Inactive") -> правило ECR (L6):
# columnCode="*" (усі колонки таблиці) + refColumnCode=<статусна колонка> + textEq <значення> -> заливка dxf.
# Рішення людини 12.10 17:35Z (квіз): «Inactive» 00CCFF додати як запис довідника й перенести CF «як в Excel».
# Переносяться ЛИШЕ правила, де порівнювана колонка - статусна колонка таблиці з довідником LAND_GENERIC_EQUIPSTATUS
# (в xlsm це колонка B). Правила на колонці C (аркуші 9a/11) НЕ доведені (що в C - невідомо) і не переносяться: лежать у
# `conditionalFormatsStatusUnproven` до відповіді замовника.
STATUS_CF_VALUES = ('Inactive',)
STATUS_REGISTRY = 'LAND_GENERIC_EQUIPSTATUS'
_status_expr = re.compile(r'^N\(\$([A-Z]{1,3})\d+\s*=\s*"([^"]+)"\)$')


def _cf_status():
    agg = {}
    unproven = {}
    for sh in out['sheets']:
        xs = next((x for x in X['sheets'] if x['name'] == sh['excelName']), None)
        if xs is None:
            continue
        for cf in xs.get('exprFormats', []):
            m = _status_expr.match((cf['f'][0] if cf['f'] else '').replace(' ', ''))
            if not m or m.group(2) not in STATUS_CF_VALUES or not cf['fill']:
                continue
            letter, value = m.group(1), m.group(2)
            for part in (cf['sqref'] or '').split():
                try:
                    c1, r1, c2, r2 = _rb(part)
                except Exception:
                    continue
                for t in sh['tables']:
                    ex = t['excel']
                    if r2 < ex['firstDataRow'] or r1 > ex['lastDataRow']:
                        continue
                    status = next((c for c in t['columns'] if c.get('excel') == letter and c.get('registry') == STATUS_REGISTRY), None)
                    if status is None:
                        k = (sh['code'], t['code'], letter, value)
                        u = unproven.setdefault(k, {'sheet': sh['code'], 'table': t['code'], 'excelColumn': letter, 'value': value,
                                                    'backgroundHex': '#' + cf['fill'], 'excelRanges': [], 'excelRuleCount': 0})
                        u['excelRanges'].append(part)
                        u['excelRuleCount'] += 1
                        continue
                    k = (sh['code'], t['code'], status['code'], value, cf['fill'])
                    a = agg.setdefault(k, {'sheet': sh['code'], 'table': t['code'], 'statusColumn': status['code'], 'excelColumn': letter,
                                           'operator': 'eq', 'value': value, 'valueKind': 'text', 'backgroundHex': '#' + cf['fill'],
                                           'foregroundHex': ('#' + cf['font']) if cf['font'] else None, 'isBold': bool(cf['bold']),
                                           'priority': int(cf['priority']) if cf.get('priority') is not None else None,
                                           'excelRanges': [], 'excelRuleCount': 0})
                    a['excelRanges'].append(part)
                    a['excelRuleCount'] += 1
    return list(agg.values()), list(unproven.values())


out['conditionalFormatsRowStatus'], out['conditionalFormatsStatusUnproven'] = _cf_status()
print('cf row-status:', len(out['conditionalFormatsRowStatus']), 'rules; unproven:', len(out['conditionalFormatsStatusUnproven']))

# ---- довідник статусів: INACTIVE (рішення людини 12.10 17:35Z). У xlsm значення «Inactive» є лише в CF-формулах
# (142 правила 00CCFF), у списку Generic_EquipStatus DropdownList!A5:A6 його немає - тож запис додається тут, у моделі,
# з якої Build створює довідник (не в сіді: довідник Land існує лише в шаблоні, сід його не створює).
_st_reg = registries.get(STATUS_REGISTRY)
if _st_reg is not None and not any(e['code'] == 'INACTIVE' for e in _st_reg['entries']):
    _st_reg['entries'].append({'code': 'INACTIVE', 'name': 'Inactive', 'en': 'Inactive', 'ru': ''})
print('styles:', len(out['styles']), [(s['code'], s['columns'], s['headers'], s['tables']) for s in out['styles']])
print('cf row-scoped:', len(out['conditionalFormatsRowScoped']))
out['registries'] = list(registries.values())
out['stats'] = {'excelSheets': len(X['sheets']), 'reportSheets': len(out['sheets']), 'tables': stat_tables, 'columns': stat_cols,
                'registries': len(registries), 'registryEntries': sum(len(r['entries']) for r in registries.values())}
# таблиці, для яких виходи відомі, а самої таблиці в шаблоні ще немає (LP-09): колонки з'являться після додавання таблиці
out['resultOutputsPending'] = {t: RESULT_TABLES[t] for t in sorted(RESULT_TABLES) if t not in RESULT_TABLES_USED}
print('result columns: tables %d, pending %s' % (len(RESULT_TABLES_USED), sorted(out['resultOutputsPending'])))
json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=None, separators=(',', ':'))
print(out['stats'])
for s in out['sheets']:
    print(s['code'], s['calc'], [(t['code'], t['layout'], t['rowMode'], len(t['columns']), len(t['rows']), len(t['rules'])) for t in s['tables']][:6])
